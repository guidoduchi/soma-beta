from __future__ import annotations

import pytest
from openpyxl import load_workbook

from soma.foundation.errors import IntegrityFailure, SomaError
from soma.foundation.identifiers import new_uuid4, utc_epoch_seconds
from soma.foundation.jobs import DurableJobCoordinator, JobTypeRegistry
from soma.foundation.persistence.uow import ReadSnapshot
from soma.product_line_sla.algorithms.cohort_state import canonical_month_bounds
from soma.product_line_sla.artifacts.xlsx_report import (
    BoundReportDestination,
    SlaReportXlsxArtifact,
)
from soma.product_line_sla.jobs import PRODUCT_LINE_SLA_JOB_CONTRACTS
from soma.product_line_sla.jobs.report_generation import SlaReportGenerationWorker
from soma.product_line_sla.repositories.reports import ReportRepository
from soma.product_line_sla.services.report_orchestration import SlaReportOrchestrationService
from soma.product_line_sla.services.report_worker import SlaReportWorkerService

MONTH = "2026-09"
DESTINATION_TOKEN = "d" * 64


def _factory(initialized_database):
    database_path, factory_builder = initialized_database
    return factory_builder(database_path)


def test_report_worker_writes_verifies_publishes_and_completes_xlsx(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    month_start, _month_end = canonical_month_bounds(MONTH)
    started = SlaReportOrchestrationService(factory).start_monthly(
        command_id=new_uuid4(),
        calendar_month=MONTH,
        as_of_utc=month_start + 1,
        customer_org_id=None,
        destination_request_token=DESTINATION_TOKEN,
    )
    report_id = str(started["report_attempt_id"])

    jobs = DurableJobCoordinator(
        factory,
        JobTypeRegistry(PRODUCT_LINE_SLA_JOB_CONTRACTS),
    )
    claim = jobs.claim_next(new_uuid4(), utc_epoch_seconds())
    assert claim is not None

    artifact = SlaReportXlsxArtifact(
        factory,
        BoundReportDestination(
            destination_request_token=DESTINATION_TOKEN,
            directory=tmp_path,
        ),
    )
    completed_id = SlaReportGenerationWorker(
        factory,
        artifact=artifact,
    ).run_to_completion(claim)
    assert completed_id == report_id

    with ReadSnapshot(factory) as snapshot:
        attempt = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert attempt is not None
        assert attempt.state == "completed"
        assert attempt.snapshot_hash is not None
        assert attempt.artifact_filename is not None
        assert attempt.artifact_sha256 is not None
        assert attempt.artifact_size_bytes is not None
        assert attempt.artifact_size_bytes > 0
        assert attempt.verified_at_utc is not None
        assert attempt.snapshot_member_count == 0
        assert attempt.snapshot_cohort_count == 0
        assert attempt.snapshot_section_row_count == 0
        final_path = tmp_path / attempt.artifact_filename
        assert final_path.is_file()
        assert final_path.stat().st_size == attempt.artifact_size_bytes

        job = snapshot.connection.execute(
            "SELECT state,claimed_run_id,claim_started_at_utc FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == ("completed", None, None)

        checkpoint = snapshot.connection.execute(
            "SELECT checkpoint_json FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()[0]
        assert '"phase":"terminal"' in str(checkpoint)

    workbook = load_workbook(final_path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "_SOMA_Metadata",
            "Members",
            "Tiers",
            "Cohorts",
            "Sections",
        ]
        metadata = dict(
            row
            for row in workbook["_SOMA_Metadata"].iter_rows(min_row=2, values_only=True)
        )
        assert metadata["contract_id"] == "SLA_REPORT_XLSX_V1"
        assert metadata["report_attempt_id"] == report_id
        assert metadata["snapshot_hash"] == attempt.snapshot_hash
        assert metadata["represented_content_sha256"] == attempt.snapshot_hash
        assert metadata["member_count"] == 0
        assert metadata["cohort_count"] == 0
        assert metadata["section_count"] == 0
    finally:
        workbook.close()


def test_sealed_snapshot_regeneration_produces_identical_candidate_bytes(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    month_start, _month_end = canonical_month_bounds(MONTH)
    started = SlaReportOrchestrationService(factory).start_monthly(
        command_id=new_uuid4(),
        calendar_month=MONTH,
        as_of_utc=month_start + 1,
        customer_org_id=None,
        destination_request_token=DESTINATION_TOKEN,
    )
    report_id = str(started["report_attempt_id"])
    jobs = DurableJobCoordinator(
        factory,
        JobTypeRegistry(PRODUCT_LINE_SLA_JOB_CONTRACTS),
    )
    claim = jobs.claim_next(new_uuid4(), utc_epoch_seconds())
    assert claim is not None

    artifact = SlaReportXlsxArtifact(
        factory,
        BoundReportDestination(DESTINATION_TOKEN, tmp_path),
    )
    worker = SlaReportGenerationWorker(factory, artifact=artifact)
    worker.snapshot_and_seal(claim)

    with ReadSnapshot(factory) as snapshot:
        attempt = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert attempt is not None
        assert attempt.state == "ready_to_generate"
        assert attempt.snapshot_hash is not None
        snapshot_hash = attempt.snapshot_hash

    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    first_path = tmp_path / candidate
    first_bytes = first_path.read_bytes()
    first_path.unlink()

    second = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
        candidate_filename=candidate,
    )
    assert second == candidate
    assert (tmp_path / second).read_bytes() == first_bytes


def test_completed_report_rejects_cancellation_and_preserves_published_artifact_t039(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    month_start, _month_end = canonical_month_bounds(MONTH)
    started = SlaReportOrchestrationService(factory).start_monthly(
        command_id=new_uuid4(),
        calendar_month=MONTH,
        as_of_utc=month_start + 1,
        customer_org_id=None,
        destination_request_token=DESTINATION_TOKEN,
    )
    report_id = str(started["report_attempt_id"])

    jobs = DurableJobCoordinator(
        factory,
        JobTypeRegistry(PRODUCT_LINE_SLA_JOB_CONTRACTS),
    )
    claim = jobs.claim_next(new_uuid4(), utc_epoch_seconds())
    assert claim is not None
    artifact = SlaReportXlsxArtifact(
        factory,
        BoundReportDestination(DESTINATION_TOKEN, tmp_path),
    )
    SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(claim)

    with ReadSnapshot(factory) as snapshot:
        attempt = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert attempt is not None
        assert attempt.state == "completed"
        assert attempt.artifact_filename is not None
        final_path = tmp_path / attempt.artifact_filename
        original_bytes = final_path.read_bytes()
        original_revision = attempt.revision
        job = snapshot.connection.execute(
            "SELECT state,claimed_run_id,claim_started_at_utc FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == ("completed", None, None)

    assert attempt.snapshot_hash is not None
    with pytest.raises(SomaError) as artifact_exc:
        artifact.write_candidate(
            report_attempt_id=report_id,
            snapshot_hash=attempt.snapshot_hash,
            destination_request_token=DESTINATION_TOKEN,
        )
    assert artifact_exc.value.code == "SLA_REPORT_ATTEMPT_STATE"

    with pytest.raises(SomaError) as excinfo:
        SlaReportOrchestrationService(factory).cancel(
            command_id=new_uuid4(),
            report_attempt_id=report_id,
            expected_attempt_revision=original_revision,
        )
    assert excinfo.value.code == "SLA_REPORT_ATTEMPT_STATE"

    with ReadSnapshot(factory) as snapshot:
        after = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert after is not None
        assert after.state == "completed"
        assert after.revision == original_revision
        job = snapshot.connection.execute(
            "SELECT state,claimed_run_id,claim_started_at_utc FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == ("completed", None, None)

    assert final_path.is_file()
    assert final_path.read_bytes() == original_bytes


def _start_sealed_report(factory, tmp_path):
    month_start, _month_end = canonical_month_bounds(MONTH)
    started = SlaReportOrchestrationService(factory).start_monthly(
        command_id=new_uuid4(),
        calendar_month=MONTH,
        as_of_utc=month_start + 1,
        customer_org_id=None,
        destination_request_token=DESTINATION_TOKEN,
    )
    report_id = str(started["report_attempt_id"])
    jobs = DurableJobCoordinator(
        factory,
        JobTypeRegistry(PRODUCT_LINE_SLA_JOB_CONTRACTS),
    )
    claim = jobs.claim_next(new_uuid4(), utc_epoch_seconds())
    assert claim is not None
    artifact = SlaReportXlsxArtifact(
        factory,
        BoundReportDestination(DESTINATION_TOKEN, tmp_path),
    )
    SlaReportGenerationWorker(factory, artifact=artifact).snapshot_and_seal(claim)
    with ReadSnapshot(factory) as snapshot:
        attempt = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert attempt is not None
        assert attempt.state == "ready_to_generate"
        assert attempt.snapshot_hash is not None
        checkpoint_json = snapshot.connection.execute(
            "SELECT checkpoint_json FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()[0]
    import json

    checkpoint = json.loads(str(checkpoint_json))
    return report_id, jobs, claim, artifact, attempt, checkpoint


def _recover_claim(jobs, claim):
    recovery_run_id = new_uuid4()
    now = utc_epoch_seconds()
    summary = jobs.recover_stale_claims(recovery_run_id, now)
    assert summary.retry_wait_count == 1
    recovered = jobs.claim_next(recovery_run_id, now)
    assert recovered is not None
    assert recovered.job_id == claim.job_id
    return recovered


def test_writing_phase_recovery_regenerates_candidate_and_completes(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, jobs, claim, artifact, attempt, checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    worker = SlaReportWorkerService(factory)
    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=attempt.snapshot_hash,
    )
    jobs.checkpoint(
        claim,
        {
            **checkpoint,
            "phase": "writing",
            "attempt_revision": int(generating["resulting_attempt_revision"]),
            "pending_command": None,
        },
    )

    recovered = _recover_claim(jobs, claim)
    assert (
        SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(
            recovered
        )
        == report_id
    )

    with ReadSnapshot(factory) as snapshot:
        completed = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert completed is not None
        assert completed.state == "completed"
        assert completed.artifact_filename is not None
        assert (tmp_path / completed.artifact_filename).is_file()


def test_verifying_phase_recovery_rebuilds_missing_exact_candidate(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, jobs, claim, artifact, attempt, checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    worker = SlaReportWorkerService(factory)
    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=attempt.snapshot_hash,
    )
    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=attempt.snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    verifying = worker.mark_verifying(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=int(generating["resulting_attempt_revision"]),
        snapshot_hash=attempt.snapshot_hash,
        candidate_filename=candidate,
    )
    jobs.checkpoint(
        claim,
        {
            **checkpoint,
            "phase": "verifying",
            "attempt_revision": int(verifying["resulting_attempt_revision"]),
            "pending_command": None,
            "candidate_filename": candidate,
        },
    )
    (tmp_path / candidate).unlink()

    recovered = _recover_claim(jobs, claim)
    assert (
        SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(
            recovered
        )
        == report_id
    )

    with ReadSnapshot(factory) as snapshot:
        completed = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert completed is not None
        assert completed.state == "completed"
        assert completed.artifact_filename is not None
        assert (tmp_path / completed.artifact_filename).is_file()


def test_publishing_recovery_without_durable_proof_fails_closed(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, jobs, claim, artifact, attempt, checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    worker = SlaReportWorkerService(factory)
    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=attempt.snapshot_hash,
    )
    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=attempt.snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    verifying = worker.mark_verifying(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=int(generating["resulting_attempt_revision"]),
        snapshot_hash=attempt.snapshot_hash,
        candidate_filename=candidate,
    )
    jobs.checkpoint(
        claim,
        {
            **checkpoint,
            "phase": "publishing",
            "attempt_revision": int(verifying["resulting_attempt_revision"]),
            "pending_command": None,
            "candidate_filename": candidate,
        },
    )
    verification = artifact.verify_candidate(
        report_attempt_id=report_id,
        snapshot_hash=attempt.snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
        candidate_filename=candidate,
    )
    proof = artifact.publish_verified(
        verification=verification,
        destination_request_token=DESTINATION_TOKEN,
    )
    final_path = tmp_path / str(proof["artifact_filename"])
    assert final_path.is_file()
    assert not (tmp_path / candidate).exists()

    recovered = _recover_claim(jobs, claim)
    assert (
        SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(
            recovered
        )
        == report_id
    )

    with ReadSnapshot(factory) as snapshot:
        failed = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert failed is not None
        assert failed.state == "failed"
        assert failed.failure_code == "SLA_REPORT_PUBLICATION_PROOF_LOST"
        assert failed.snapshot_hash is None
        assert ReportRepository.counts(snapshot.connection, report_id) == (0, 0, 0, 0)
        job = snapshot.connection.execute(
            "SELECT state,last_error_code FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == (
            "failed",
            "SLA_REPORT_PUBLICATION_PROOF_LOST",
        )

    # Atomic publication bytes may exist, but without durable proof they are
    # external housekeeping only and cancellation/failure never deletes them.
    assert final_path.is_file()



def test_f021_partial_candidate_is_never_verified_and_can_regenerate_from_sealed_snapshot(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, _jobs, _claim, artifact, attempt, _checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    assert attempt.snapshot_hash is not None
    snapshot_hash = attempt.snapshot_hash
    worker = SlaReportWorkerService(factory)

    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=snapshot_hash,
    )
    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    verifying = worker.mark_verifying(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=int(generating["resulting_attempt_revision"]),
        snapshot_hash=snapshot_hash,
        candidate_filename=candidate,
    )

    candidate_path = tmp_path / candidate
    original = candidate_path.read_bytes()
    assert len(original) > 128
    candidate_path.write_bytes(original[:64])

    with pytest.raises(SomaError) as failure:
        artifact.verify_candidate(
            report_attempt_id=report_id,
            snapshot_hash=snapshot_hash,
            destination_request_token=DESTINATION_TOKEN,
            candidate_filename=candidate,
        )
    assert failure.value.code == "SLA_REPORT_ARTIFACT_VERIFY_FAILED"

    with ReadSnapshot(factory) as snapshot:
        current = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert current is not None
        assert current.state == "verifying"
        assert current.revision == int(verifying["resulting_attempt_revision"])
        assert current.artifact_filename == candidate
        assert current.artifact_sha256 is None
        assert current.verified_at_utc is None
        assert current.completed_at_utc is None

    regenerated = artifact.recover_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
        candidate_filename=candidate,
    )
    assert regenerated == candidate
    verification = artifact.verify_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
        candidate_filename=candidate,
    )
    assert verification.report_attempt_id == report_id
    assert verification.snapshot_hash == snapshot_hash


def test_f022_reopenable_workbook_with_wrong_shape_cannot_verify_or_complete(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, _jobs, _claim, artifact, attempt, _checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    assert attempt.snapshot_hash is not None
    snapshot_hash = attempt.snapshot_hash
    worker = SlaReportWorkerService(factory)

    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=snapshot_hash,
    )
    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    verifying = worker.mark_verifying(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=int(generating["resulting_attempt_revision"]),
        snapshot_hash=snapshot_hash,
        candidate_filename=candidate,
    )

    path = tmp_path / candidate
    workbook = load_workbook(path)
    try:
        workbook["Members"]["A1"] = "corrupted_member_header"
        workbook.save(path)
    finally:
        workbook.close()

    with pytest.raises(IntegrityFailure, match="Members header is invalid"):
        artifact.verify_candidate(
            report_attempt_id=report_id,
            snapshot_hash=snapshot_hash,
            destination_request_token=DESTINATION_TOKEN,
            candidate_filename=candidate,
        )

    with ReadSnapshot(factory) as snapshot:
        current = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert current is not None
        assert current.state == "verifying"
        assert current.revision == int(verifying["resulting_attempt_revision"])
        assert current.artifact_filename == candidate
        assert current.artifact_sha256 is None
        assert current.artifact_size_bytes is None
        assert current.verified_at_utc is None
        assert current.completed_at_utc is None


def test_f028_missing_external_artifact_does_not_mutate_completed_database_evidence(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    month_start, _month_end = canonical_month_bounds(MONTH)
    started = SlaReportOrchestrationService(factory).start_monthly(
        command_id=new_uuid4(),
        calendar_month=MONTH,
        as_of_utc=month_start + 1,
        customer_org_id=None,
        destination_request_token=DESTINATION_TOKEN,
    )
    report_id = str(started["report_attempt_id"])
    jobs = DurableJobCoordinator(
        factory,
        JobTypeRegistry(PRODUCT_LINE_SLA_JOB_CONTRACTS),
    )
    claim = jobs.claim_next(new_uuid4(), utc_epoch_seconds())
    assert claim is not None
    artifact = SlaReportXlsxArtifact(
        factory,
        BoundReportDestination(DESTINATION_TOKEN, tmp_path),
    )
    SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(claim)

    with ReadSnapshot(factory) as snapshot:
        before = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert before is not None
        assert before.state == "completed"
        assert before.artifact_filename is not None
        snapshot_hash = ReportRepository.snapshot_hash(
            snapshot.connection,
            report_id,
        )
        semantic = ReportRepository.snapshot_semantic_value(
            snapshot.connection,
            report_id,
        )
        final_path = tmp_path / before.artifact_filename
        assert final_path.is_file()

    final_path.unlink()
    assert not final_path.exists()

    with ReadSnapshot(factory) as snapshot:
        after = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert after == before
        assert after.state == "completed"
        assert ReportRepository.snapshot_hash(
            snapshot.connection,
            report_id,
        ) == snapshot_hash
        assert ReportRepository.snapshot_semantic_value(
            snapshot.connection,
            report_id,
        ) == semantic
        job = snapshot.connection.execute(
            "SELECT state FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == ("completed",)



def test_f020_recovery_from_exact_sealed_checkpoint_regenerates_without_live_requery(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, jobs, claim, artifact, attempt, _checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    assert attempt.state == "ready_to_generate"
    assert attempt.snapshot_hash is not None
    sealed_hash = attempt.snapshot_hash

    recovered = _recover_claim(jobs, claim)
    assert (
        SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(
            recovered
        )
        == report_id
    )

    with ReadSnapshot(factory) as snapshot:
        completed = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert completed is not None
        assert completed.state == "completed"
        assert completed.snapshot_hash == sealed_hash
        assert completed.artifact_filename is not None
        assert (tmp_path / completed.artifact_filename).is_file()
        job = snapshot.connection.execute(
            "SELECT state FROM durable_jobs WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == ("completed",)


def test_f026_cancellation_wins_before_complete_and_published_bytes_stay_non_authoritative(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, jobs, claim, artifact, attempt, checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    assert attempt.snapshot_hash is not None
    snapshot_hash = attempt.snapshot_hash
    worker = SlaReportWorkerService(factory)

    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=snapshot_hash,
    )
    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    verifying = worker.mark_verifying(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=int(generating["resulting_attempt_revision"]),
        snapshot_hash=snapshot_hash,
        candidate_filename=candidate,
    )
    jobs.checkpoint(
        claim,
        {
            **checkpoint,
            "phase": "verifying",
            "attempt_revision": int(verifying["resulting_attempt_revision"]),
            "pending_command": None,
            "candidate_filename": candidate,
        },
    )
    verification = artifact.verify_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
        candidate_filename=candidate,
    )

    # Model bytes physically landing while SQLite completion has not committed.
    proof = artifact.publish_verified(
        verification=verification,
        destination_request_token=DESTINATION_TOKEN,
    )
    final_path = tmp_path / str(proof["artifact_filename"])
    assert final_path.is_file()
    assert not (tmp_path / candidate).exists()

    cancel_command = new_uuid4()
    cancelled = SlaReportOrchestrationService(factory).cancel(
        command_id=cancel_command,
        report_attempt_id=report_id,
        expected_attempt_revision=int(verifying["resulting_attempt_revision"]),
    )
    assert cancelled["state"] == "cancelled"
    assert cancelled["snapshot_hash"] is None
    assert cancelled["artifact_filename"] is None

    complete_command = new_uuid4()
    with pytest.raises(SomaError) as completion_lost:
        worker.complete(
            command_id=complete_command,
            report_attempt_id=report_id,
            expected_attempt_revision=int(verifying["resulting_attempt_revision"]),
            snapshot_hash=snapshot_hash,
            artifact_completion=proof,
        )
    assert completion_lost.value.code == "SLA_REPORT_ATTEMPT_STATE"

    with ReadSnapshot(factory) as snapshot:
        final = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert final is not None
        assert final.state == "cancelled"
        assert final.snapshot_hash is None
        assert final.artifact_filename is None
        assert final.artifact_sha256 is None
        assert final.completed_at_utc is not None
        assert ReportRepository.counts(snapshot.connection, report_id) == (0, 0, 0, 0)
        job = snapshot.connection.execute(
            "SELECT state,claimed_run_id,claim_started_at_utc FROM durable_jobs "
            "WHERE job_id=?",
            (claim.job_id,),
        ).fetchone()
        assert tuple(job) == ("cancelled", None, None)
        assert snapshot.connection.execute(
            "SELECT COUNT(*) FROM command_receipts WHERE command_id=?",
            (complete_command,),
        ).fetchone()[0] == 0

    # The filesystem is outside SQLite authority. These already-published bytes
    # remain housekeeping only; they do not resurrect completed report evidence.
    assert final_path.is_file()



def test_f033_recovered_exact_candidate_reconstruction_keeps_verifying_authority_unchanged(
    initialized_database,
    tmp_path,
) -> None:
    factory = _factory(initialized_database)
    report_id, jobs, claim, artifact, attempt, checkpoint = _start_sealed_report(
        factory,
        tmp_path,
    )
    assert attempt.snapshot_hash is not None
    snapshot_hash = attempt.snapshot_hash
    worker = SlaReportWorkerService(factory)

    generating = worker.mark_generating(
        command_id=new_uuid4(),
        report_attempt_id=report_id,
        expected_attempt_revision=attempt.revision,
        snapshot_hash=snapshot_hash,
    )
    candidate = artifact.write_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
    )
    verify_command = new_uuid4()
    verifying = worker.mark_verifying(
        command_id=verify_command,
        report_attempt_id=report_id,
        expected_attempt_revision=int(generating["resulting_attempt_revision"]),
        snapshot_hash=snapshot_hash,
        candidate_filename=candidate,
    )
    verifying_revision = int(verifying["resulting_attempt_revision"])
    jobs.checkpoint(
        claim,
        {
            **checkpoint,
            "phase": "verifying",
            "attempt_revision": verifying_revision,
            "pending_command": None,
            "candidate_filename": candidate,
        },
    )
    candidate_path = tmp_path / candidate
    assert candidate_path.is_file()
    candidate_path.unlink()

    recovered = _recover_claim(jobs, claim)

    rebuilt = artifact.recover_candidate(
        report_attempt_id=report_id,
        snapshot_hash=snapshot_hash,
        destination_request_token=DESTINATION_TOKEN,
        candidate_filename=candidate,
    )
    assert rebuilt == candidate
    assert candidate_path.is_file()

    with ReadSnapshot(factory) as snapshot:
        unchanged = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert unchanged is not None
        assert unchanged.state == "verifying"
        assert unchanged.revision == verifying_revision
        assert unchanged.snapshot_hash == snapshot_hash
        assert unchanged.artifact_filename == candidate
        assert unchanged.artifact_sha256 is None
        assert unchanged.artifact_size_bytes is None
        assert unchanged.verified_at_utc is None
        assert unchanged.completed_at_utc is None
        assert snapshot.connection.execute(
            "SELECT COUNT(*) FROM command_receipts WHERE command_id=?",
            (verify_command,),
        ).fetchone()[0] == 1
        assert snapshot.connection.execute(
            "SELECT COUNT(*) FROM command_receipt_results WHERE command_id=?",
            (verify_command,),
        ).fetchone()[0] == 1
        assert snapshot.connection.execute(
            "SELECT COUNT(*) FROM audit_events "
            "WHERE command_id=? AND action_type='sla.report.progressed'",
            (verify_command,),
        ).fetchone()[0] == 1

    temp_candidates = sorted(
        path.name
        for path in tmp_path.iterdir()
        if path.is_file() and path.name.startswith(".") and path.suffix == ".xlsx"
    )
    assert temp_candidates == [candidate]

    assert (
        SlaReportGenerationWorker(factory, artifact=artifact).run_to_completion(
            recovered
        )
        == report_id
    )
    with ReadSnapshot(factory) as snapshot:
        completed = ReportRepository.get_attempt(snapshot.connection, report_id)
        assert completed is not None
        assert completed.state == "completed"
        assert completed.snapshot_hash == snapshot_hash
        assert snapshot.connection.execute(
            "SELECT COUNT(*) FROM audit_events "
            "WHERE command_id=? AND action_type='sla.report.progressed'",
            (verify_command,),
        ).fetchone()[0] == 1
