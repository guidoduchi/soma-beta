from __future__ import annotations

from io import BytesIO
from uuid import UUID
import zipfile

import pytest
from openpyxl import load_workbook

from soma.foundation.errors import ValidationError
from soma.infrastructure.contracts.infrastructure import validate_value
from soma.infrastructure.jobs.infrastructure_workbooks import (
    INFRASTRUCTURE_XLSX_LIMITS, preflight_infrastructure_workbook,
    verify_profile_workbook, write_profile_workbook,
)
from soma.infrastructure.domain.workbooks import (
    HEADERS, METADATA_KEYS, SHEET_ORDER, artifact_filename,
)


def test_infrastructure_workbook_profile_and_local_filename():
    assert SHEET_ORDER == ("Metadata", "Network Elements", "IP Addresses")
    assert HEADERS["Metadata"] == ("Key", "Value")
    assert len(HEADERS["Network Elements"]) == 22
    assert len(HEADERS["IP Addresses"]) == 5
    assert METADATA_KEYS == (
        "FormatId", "WorkbookVersion", "Mode", "SourceInstallationScopeId",
        "GeneratedAtUtc", "ExportScopeJson",
    )
    assert artifact_filename(
        "round_trip", "12345678-1234-4234-8234-123456789abc", 0,
    ) == "SOMA_Infrastructure_round_trip_19691231_190000_12345678.xlsx"
    with pytest.raises(ValidationError):
        artifact_filename("other", "12345678-1234-4234-8234-123456789abc", 0)
    with pytest.raises(ValidationError):
        artifact_filename("discovery", "not-a-uuid", 0)


def test_profile_writer_emits_fixed_sheets_and_literal_formula_like_text():
    element = [None] * len(HEADERS["Network Elements"])
    element[1] = "=1+1"
    element[6] = "Site"
    element[7] = "1 Main Street"
    stream = BytesIO()
    counts = write_profile_workbook(
        stream, mode="discovery", scope={"scope_kind": "all"},
        data_instance_id="12345678-1234-4234-8234-123456789abc",
        generated_at_utc=0, network_elements=[element],
        ip_addresses=[(None, None, "=1+1", "192.0.2.1", True)],
    )
    assert counts == (1, 1)
    digest, size = verify_profile_workbook(
        stream, mode="discovery",
        data_instance_id="12345678-1234-4234-8234-123456789abc",
        expected_scope={"scope_kind": "all"}, expected_generated_at_utc=0,
        expected_network_rows=1, expected_ip_rows=1,
    )
    assert len(digest) == 64 and size == len(stream.getvalue())
    stream.seek(0)
    workbook = load_workbook(stream, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == list(SHEET_ORDER)
        for name in SHEET_ORDER:
            assert tuple(cell.value for cell in next(workbook[name].rows)) == HEADERS[name]
        metadata = dict(workbook["Metadata"].values)
        assert metadata["FormatId"] == "SOMA-INFRA-XLSX"
        assert metadata["WorkbookVersion"] == "1.0"
        assert metadata["GeneratedAtUtc"] == "1970-01-01T00:00:00Z"
        value = workbook["Network Elements"]["B2"]
        assert value.value == "=1+1" and value.data_type == "s"
        ip_value = workbook["IP Addresses"]["C2"]
        assert ip_value.value == "=1+1" and ip_value.data_type == "s"
    finally:
        workbook.close()


def test_profile_writer_rejects_scope_shape_before_publication():
    with pytest.raises(ValidationError):
        write_profile_workbook(
            BytesIO(), mode="registration_template",
            scope={"scope_kind": "all", "site_id": "12345678-1234-4234-8234-123456789abc"},
            data_instance_id="12345678-1234-4234-8234-123456789abc",
            generated_at_utc=0,
        )


def test_explicit_element_scope_limit_fits_metadata_cell():
    ids = [str(UUID(int=index + 1, version=4)) for index in range(400)]
    scope = {"scope_kind": "network_elements", "network_element_ids": ids}
    assert len(validate_value("INFRA_EXPORT_SCOPE_V1", scope)["network_element_ids"]) == 400
    stream = BytesIO()
    write_profile_workbook(
        stream, mode="round_trip", scope=scope,
        data_instance_id="12345678-1234-4234-8234-123456789abc",
        generated_at_utc=0,
    )
    stream.seek(0)
    workbook = load_workbook(stream, read_only=True)
    try:
        metadata = dict(workbook["Metadata"].values)
        assert len(metadata["ExportScopeJson"].encode("utf-8")) <= 16_384
    finally:
        workbook.close()
    with pytest.raises(ValidationError):
        validate_value("INFRA_EXPORT_SCOPE_V1", {
            "scope_kind": "network_elements", "network_element_ids": ids + [ids[0]],
        })


def test_profile_verifier_rejects_wrong_row_count():
    stream = BytesIO()
    write_profile_workbook(
        stream, mode="registration_template", scope={"scope_kind": "all"},
        data_instance_id="12345678-1234-4234-8234-123456789abc",
        generated_at_utc=0,
    )
    with pytest.raises(ValidationError):
        verify_profile_workbook(
            stream, mode="registration_template",
            data_instance_id="12345678-1234-4234-8234-123456789abc",
            expected_scope={"scope_kind": "all"}, expected_generated_at_utc=0,
            expected_network_rows=1, expected_ip_rows=0,
        )


def test_profile_verifier_rejects_metadata_from_different_generation():
    stream = BytesIO()
    write_profile_workbook(
        stream, mode="registration_template", scope={"scope_kind": "all"},
        data_instance_id="12345678-1234-4234-8234-123456789abc",
        generated_at_utc=0,
    )
    with pytest.raises(ValidationError):
        verify_profile_workbook(
            stream, mode="registration_template",
            data_instance_id="12345678-1234-4234-8234-123456789abc",
            expected_scope={"scope_kind": "all"}, expected_generated_at_utc=1,
            expected_network_rows=0, expected_ip_rows=0,
        )


def test_profile_verifier_rejects_duplicate_archive_part():
    stream = BytesIO()
    write_profile_workbook(
        stream, mode="registration_template", scope={"scope_kind": "all"},
        data_instance_id="12345678-1234-4234-8234-123456789abc",
        generated_at_utc=0,
    )
    with pytest.warns(UserWarning):
        with zipfile.ZipFile(stream, "a") as archive:
            archive.writestr("[Content_Types].xml", b"<Types/>")
    with pytest.raises(ValidationError):
        verify_profile_workbook(
            stream, mode="registration_template",
            data_instance_id="12345678-1234-4234-8234-123456789abc",
            expected_scope={"scope_kind": "all"}, expected_generated_at_utc=0,
            expected_network_rows=0, expected_ip_rows=0,
        )


def test_import_preflight_uses_infrastructure_limits_and_same_captured_bytes(tmp_path, monkeypatch):
    from soma.ticket_import.parsing import xlsx_security

    source = tmp_path / "source.xlsx"
    with source.open("wb") as stream:
        write_profile_workbook(
            stream, mode="registration_template", scope={"scope_kind": "all"},
            data_instance_id="12345678-1234-4234-8234-123456789abc",
            generated_at_utc=0,
        )
    assert INFRASTRUCTURE_XLSX_LIMITS.zip_entries == 20_000
    monkeypatch.setattr(xlsx_security, "MAX_ZIP_ENTRIES", 4)
    preflight = preflight_infrastructure_workbook(str(source))
    try:
        assert preflight.entry_count == 11
        assert preflight.semantic_stream().read() == source.read_bytes()
    finally:
        preflight.close()


def test_import_preflight_rejects_unsafe_part_without_modifying_source(tmp_path):
    from soma.foundation.errors import SomaError

    source = tmp_path / "unsafe.xlsx"
    with source.open("wb") as stream:
        write_profile_workbook(
            stream, mode="registration_template", scope={"scope_kind": "all"},
            data_instance_id="12345678-1234-4234-8234-123456789abc",
            generated_at_utc=0,
        )
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("../escape.xml", b"<escape/>")
    before = source.read_bytes()
    with pytest.raises(SomaError) as error:
        preflight_infrastructure_workbook(str(source))
    assert error.value.code == "WORKBOOK_UNSAFE"
    assert source.read_bytes() == before
