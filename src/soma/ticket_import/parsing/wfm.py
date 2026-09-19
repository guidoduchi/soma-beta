from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import load_workbook

from soma.foundation.errors import SomaError, ValidationError
from soma.objectives_tasks.services.task_planning import validate_wfm_task_no
from soma.ticket_import.profiles.registry import require_field_spec, require_profile_versions
from soma.ticket_import.reconciliation.engine import LogicalField, LogicalRow, row_logical_sha256
from soma.ticket_import.repositories.observations import NormalizedObservationEvidence
from soma.tickets.validation import validate_rfc_no

from .advanced_search import (
    ParsedFinding,
    _MAX_CELL_UTF8_BYTES,
    _MAX_DEFINED_NAMES,
    _MAX_HEADER_SEARCH_NONEMPTY_ROWS,
    _MAX_LOGICAL_CELLS_TOTAL,
    _MAX_MATRIX_ROWS,
    _MAX_PHYSICAL_COLUMNS,
    _MAX_SHARED_STRINGS,
    _MAX_WORKSHEETS,
    _Matrix,
    _blank_field,
    _controlled_key,
    _finding,
    _header_key,
    _is_formula,
    _logical_to_normalized,
    _malformed_field,
    _parse_instant_field,
    _parse_text_field,
    _row_is_physically_nonempty,
    _scalar_source_text,
    _source_error,
    _trim_spreadsheet_whitespace,
    _validate_source_text,
)
from .xlsx_security import XlsxPreflightResult

_SOURCE_FAMILY = "wfm_service_provider"
_HEADER_ALIASES = {
    "rfc_no": ("RFC No.", "RFC No"),
    "rfc_status": ("RFC Status",),
    "task_no": ("Task No.", "Task No"),
    "task_status": ("Task Status",),
    "dispatch_progress": ("Dispatch Progress",),
    "planned_start": ("Planned Start Time",),
    "planned_end": ("Planned End Time",),
    "rep_office": ("Rep Office",),
    "risk_level": ("Risk Level",),
    "task_name": ("Task Name",),
    "description": ("Description",),
    "customer_organization": ("Customer Organization",),
}
_RFC_STATUS_VALUES = ("Implement", "Closed", "Cancelled")
_TASK_STATUS_VALUES = ("Complete", "Plan Cancel")


@dataclass(frozen=True, slots=True)
class ParsedWfmRow:
    observation: NormalizedObservationEvidence
    findings: tuple[ParsedFinding, ...]
    logical_row: LogicalRow
    classification: str
    duplicate_count: int
    conflict_variant_count: int


@dataclass(frozen=True, slots=True)
class WfmParseResult:
    source_family: str
    source_profile_id: str
    header_registry_id: str
    vocabulary_registry_id: str
    parser_profile_id: str
    source_chronology_utc: int
    rows: tuple[ParsedWfmRow, ...]
    global_findings: tuple[ParsedFinding, ...]

    @property
    def observations(self) -> tuple[NormalizedObservationEvidence, ...]:
        return tuple(row.observation for row in self.rows)


def _build_header_registry() -> dict[str, str]:
    output: dict[str, str] = {}
    for field_key, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            key = _header_key(alias)
            assert key is not None
            previous = output.setdefault(key, field_key)
            if previous != field_key:
                raise RuntimeError("WFM header registry aliases collide")
    return output


_HEADER_KEYS = _build_header_registry()
_RFC_STATUS_REGISTRY = {_controlled_key(value): value for value in _RFC_STATUS_VALUES}
_TASK_STATUS_REGISTRY = {_controlled_key(value): value for value in _TASK_STATUS_VALUES}


def _resolve_header_row(row, *, sheet_ordinal: int, row_ordinal: int) -> dict[str, int] | None:
    resolved: dict[str, int] = {}
    for column_ordinal, cell in enumerate(row, start=1):
        if column_ordinal > _MAX_PHYSICAL_COLUMNS:
            raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
        key = _header_key(cell.value)
        if key is None:
            continue
        field_key = _HEADER_KEYS.get(key)
        if field_key is None:
            continue
        if field_key in resolved:
            raise _source_error(
                "SOURCE_DUPLICATE_SEMANTIC_HEADER",
                f"worksheet {sheet_ordinal} row {row_ordinal} repeats semantic header {field_key}",
            )
        resolved[field_key] = column_ordinal
    return resolved if "rfc_no" in resolved or "task_no" in resolved else None


def _discover_matrix(workbook) -> _Matrix:
    if len(workbook.worksheets) > _MAX_WORKSHEETS:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the worksheet-count ceiling")
    if len(getattr(workbook, "shared_strings", ())) > _MAX_SHARED_STRINGS:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the shared-string ceiling")
    if len(getattr(workbook, "defined_names", ())) > _MAX_DEFINED_NAMES:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the defined-name ceiling")

    candidates: list[_Matrix] = []
    scanned_cells = 0
    for sheet_ordinal, worksheet in enumerate(workbook.worksheets, start=1):
        nonempty_seen = 0
        for row in worksheet.iter_rows():
            if len(row) > _MAX_PHYSICAL_COLUMNS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                scanned_cells += 1
                if scanned_cells > _MAX_LOGICAL_CELLS_TOTAL:
                    raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the logical-cell ceiling")
                if isinstance(value, str):
                    if "\x00" in value:
                        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "workbook cell text contains NUL")
                    if len(value.encode("utf-8", errors="strict")) > _MAX_CELL_UTF8_BYTES:
                        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook cell exceeds the UTF-8 byte ceiling")
            if not _row_is_physically_nonempty(row):
                continue
            nonempty_seen += 1
            if nonempty_seen > _MAX_HEADER_SEARCH_NONEMPTY_ROWS:
                continue
            row_ordinal = int(row[0].row) if row else nonempty_seen
            columns = _resolve_header_row(row, sheet_ordinal=sheet_ordinal, row_ordinal=row_ordinal)
            if columns is not None:
                candidates.append(_Matrix(sheet_ordinal, worksheet.title, row_ordinal, columns))
    if not candidates:
        raise _source_error("SOURCE_MATRIX_NOT_FOUND", "no Service Provider WFM matrix was found")
    if len(candidates) != 1:
        raise _source_error("SOURCE_MATRIX_AMBIGUOUS", "multiple Service Provider WFM matrices were found")
    matrix = candidates[0]
    missing = [field for field in ("rfc_no", "task_no") if field not in matrix.columns]
    if missing:
        raise _source_error(
            "SOURCE_REQUIRED_HEADER_MISSING",
            f"Service Provider WFM required identity header {missing[0]} is absent",
        )
    return matrix


def _global_coverage_findings(matrix: _Matrix) -> tuple[ParsedFinding, ...]:
    return tuple(
        _finding(
            "SOURCE_HEADER_COVERAGE_MISSING",
            "warning",
            "workbook",
            f"registered Service Provider WFM header {field_key} is absent",
            field_key=field_key,
        )
        for field_key in sorted(_HEADER_ALIASES)
        if field_key not in {"rfc_no", "task_no"} and field_key not in matrix.columns
    )


def _identity_text(value: object) -> str | None:
    source_text = _scalar_source_text(value)
    if source_text is None or isinstance(value, bool):
        return None
    candidate = _trim_spreadsheet_whitespace(source_text)
    return candidate or None


def _valid_rfc(value: object) -> str | None:
    candidate = _identity_text(value)
    if candidate is None:
        return None
    try:
        return validate_rfc_no(candidate)
    except SomaError:
        return None


def _valid_task(value: object) -> str | None:
    candidate = _identity_text(value)
    if candidate is None:
        return None
    try:
        return validate_wfm_task_no(candidate)
    except SomaError:
        return None


def _parse_controlled(spec, value: object, *, sheet_ordinal: int, row_ordinal: int):
    source_text = _scalar_source_text(value)
    _validate_source_text(source_text, spec)
    if source_text is None:
        return _blank_field(spec), ()
    if not isinstance(value, (str, int, float)) or isinstance(value, bool):
        return _malformed_field(spec, source_text), ()
    trimmed = _trim_spreadsheet_whitespace(source_text)
    if not trimmed:
        return _blank_field(spec, source_text=source_text), ()
    registry = _RFC_STATUS_REGISTRY if spec.field_key == "rfc_status" else _TASK_STATUS_REGISTRY
    canonical = registry.get(_controlled_key(trimmed))
    if canonical is None:
        field = LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="unknown",
            value_kind="controlled",
            vocabulary_id=spec.vocabulary_id,
            source_text=source_text,
        )
        return field, (
            _finding(
                "SOURCE_CONTROLLED_VALUE_UNKNOWN",
                "warning",
                "field",
                f"{spec.field_key} is outside its accepted controlled vocabulary",
                field_key=spec.field_key,
                sheet_ordinal=sheet_ordinal,
                row_ordinal=row_ordinal,
            ),
        )
    return (
        LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="usable",
            value_kind="controlled",
            vocabulary_id=spec.vocabulary_id,
            source_text=source_text,
            normalized_text=canonical,
        ),
        (),
    )


def _parse_field(field_key: str, cell, *, sheet_ordinal: int, row_ordinal: int):
    spec = require_field_spec(_SOURCE_FAMILY, field_key)
    if _is_formula(cell):
        raise _source_error(
            "SOURCE_FORMULA_IN_SEMANTIC_FIELD",
            f"registered semantic field {field_key} contains a formula",
        )
    if spec.value_kind == "text":
        return _parse_text_field(spec, cell.value), ()
    if spec.value_kind == "controlled":
        return _parse_controlled(spec, cell.value, sheet_ordinal=sheet_ordinal, row_ordinal=row_ordinal)
    if spec.value_kind == "instant":
        return _parse_instant_field(spec, cell.value), ()
    raise RuntimeError("unsupported Service Provider WFM field kind")


def _plan_findings(
    fields: list[LogicalField],
    *,
    sheet_ordinal: int,
    row_ordinal: int,
) -> tuple[ParsedFinding, ...]:
    by_key = {field.field_key: field for field in fields}
    start = by_key.get("planned_start")
    end = by_key.get("planned_end")
    if start is None or end is None:
        return ()
    if start.value_state == "blank" and end.value_state == "blank":
        return ()
    if (start.value_state == "blank") != (end.value_state == "blank"):
        return (
            _finding(
                "WFM_PLAN_INCOMPLETE",
                "error",
                "row",
                "WFM source plan supplies only one planned endpoint",
                field_key="planned_start" if start.value_state == "blank" else "planned_end",
                sheet_ordinal=sheet_ordinal,
                row_ordinal=row_ordinal,
            ),
        )
    if start.value_state != "usable" or end.value_state != "usable":
        return (
            _finding(
                "WFM_PLAN_INVALID_INTERVAL",
                "error",
                "row",
                "WFM source plan endpoints are not both usable instants",
                sheet_ordinal=sheet_ordinal,
                row_ordinal=row_ordinal,
            ),
        )
    assert start.integer_value is not None and end.integer_value is not None
    if end.integer_value <= start.integer_value:
        return (
            _finding(
                "WFM_PLAN_INVALID_INTERVAL",
                "error",
                "row",
                "WFM source plan end must be strictly later than start",
                sheet_ordinal=sheet_ordinal,
                row_ordinal=row_ordinal,
            ),
        )
    return ()


def _classify_rows(rows: list[ParsedWfmRow]) -> list[ParsedWfmRow]:
    groups: dict[str, dict[str, list[ParsedWfmRow]]] = {}
    for row in rows:
        task_no = row.logical_row.canonical_primary_id
        if row.logical_row.identity_state != "valid" or task_no is None:
            continue
        groups.setdefault(task_no, {}).setdefault(row.observation.row_logical_sha256, []).append(row)

    classification: dict[tuple[str, str], tuple[str, int, int]] = {}
    for task_no, variants in groups.items():
        variant_count = len(variants)
        for digest, occurrences in variants.items():
            count = len(occurrences)
            classification[(task_no, digest)] = (
                "CONFLICT_MEMBER" if variant_count > 1 else "EQUIVALENT_DUPLICATE" if count > 1 else "UNIQUE",
                count,
                variant_count,
            )

    output: list[ParsedWfmRow] = []
    for parsed in rows:
        task_no = parsed.logical_row.canonical_primary_id
        if parsed.logical_row.identity_state != "valid" or task_no is None:
            output.append(parsed)
            continue
        row_class, duplicate_count, conflict_count = classification[(task_no, parsed.observation.row_logical_sha256)]
        findings = list(parsed.findings)
        if row_class == "EQUIVALENT_DUPLICATE":
            findings.append(
                _finding(
                    "SOURCE_EQUIVALENT_DUPLICATE",
                    "warning",
                    "row",
                    "equivalent WFM rows share one logical proposal variant",
                    sheet_ordinal=parsed.observation.sheet_ordinal,
                    row_ordinal=parsed.observation.row_ordinal,
                )
            )
        elif row_class == "CONFLICT_MEMBER":
            findings.append(
                _finding(
                    "SOURCE_IDENTITY_CONFLICT",
                    "error",
                    "identity",
                    "same WFM Task No has conflicting owning RFC or logical source content",
                    sheet_ordinal=parsed.observation.sheet_ordinal,
                    row_ordinal=parsed.observation.row_ordinal,
                )
            )
        output.append(
            replace(
                parsed,
                findings=tuple(findings),
                classification=row_class,
                duplicate_count=duplicate_count,
                conflict_variant_count=conflict_count,
            )
        )
    return output


def parse_wfm_service_provider(
    path: Path,
    *,
    preflight: XlsxPreflightResult,
    source_chronology_utc: int,
) -> WfmParseResult:
    if type(source_chronology_utc) is not int or source_chronology_utc < 0:
        raise ValidationError("source_chronology_utc must be a nonnegative whole-second epoch")
    requested = Path(path)
    try:
        requested_resolved = requested.resolve(strict=True)
        preflight_resolved = Path(preflight.path).resolve(strict=True)
    except OSError as exc:
        raise _source_error("IMPORT_SOURCE_UNAVAILABLE", "preflighted workbook is no longer available") from exc
    if requested_resolved != preflight_resolved:
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "parser path does not match preflighted workbook")

    versions = require_profile_versions(_SOURCE_FAMILY)
    try:
        workbook = load_workbook(requested_resolved, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "preflighted workbook could not be opened semantically") from exc

    try:
        matrix = _discover_matrix(workbook)
        global_findings = _global_coverage_findings(matrix)
        worksheet = workbook.worksheets[matrix.sheet_ordinal - 1]
        rows: list[ParsedWfmRow] = []
        logical_cells = 0
        matrix_rows = 0
        for physical in worksheet.iter_rows(min_row=matrix.header_row_ordinal + 1):
            if len(physical) > _MAX_PHYSICAL_COLUMNS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
            if not _row_is_physically_nonempty(physical):
                continue
            matrix_rows += 1
            if matrix_rows > _MAX_MATRIX_ROWS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "Service Provider WFM matrix exceeds the row ceiling")
            logical_cells += len(matrix.columns)
            if logical_cells > _MAX_LOGICAL_CELLS_TOTAL:
                raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the logical-cell ceiling")
            row_ordinal = int(physical[0].row) if physical else matrix.header_row_ordinal + matrix_rows
            cells = {index + 1: cell for index, cell in enumerate(physical)}

            rfc_cell = cells.get(matrix.columns["rfc_no"])
            task_cell = cells.get(matrix.columns["task_no"])
            if rfc_cell is not None and _is_formula(rfc_cell):
                raise _source_error("SOURCE_FORMULA_IN_SEMANTIC_FIELD", "registered semantic field rfc_no contains a formula")
            if task_cell is not None and _is_formula(task_cell):
                raise _source_error("SOURCE_FORMULA_IN_SEMANTIC_FIELD", "registered semantic field task_no contains a formula")
            rfc_no = _valid_rfc(None if rfc_cell is None else rfc_cell.value)
            task_no = _valid_task(None if task_cell is None else task_cell.value)
            row_findings: list[ParsedFinding] = []
            if rfc_no is None:
                row_findings.append(
                    _finding(
                        "WFM_PARENT_RFC_INVALID",
                        "error",
                        "identity",
                        "WFM RFC No is not exact NC followed by fourteen ASCII digits",
                        field_key="rfc_no",
                        sheet_ordinal=matrix.sheet_ordinal,
                        row_ordinal=row_ordinal,
                    )
                )
            if task_no is None:
                row_findings.append(
                    _finding(
                        "WFM_TASK_ID_INVALID",
                        "error",
                        "identity",
                        "WFM Task No is not exact TK followed by fourteen ASCII digits",
                        field_key="task_no",
                        sheet_ordinal=matrix.sheet_ordinal,
                        row_ordinal=row_ordinal,
                    )
                )

            logical_fields: list[LogicalField] = []
            for field_key, column_ordinal in matrix.columns.items():
                if field_key in {"rfc_no", "task_no"}:
                    continue
                cell = cells.get(column_ordinal)
                if cell is None:
                    class _Blank:
                        value = None
                        data_type = "n"
                    cell = _Blank()
                field, findings = _parse_field(
                    field_key,
                    cell,
                    sheet_ordinal=matrix.sheet_ordinal,
                    row_ordinal=row_ordinal,
                )
                logical_fields.append(field)
                row_findings.extend(findings)
            row_findings.extend(
                _plan_findings(
                    logical_fields,
                    sheet_ordinal=matrix.sheet_ordinal,
                    row_ordinal=row_ordinal,
                )
            )

            valid = rfc_no is not None and task_no is not None
            logical_row = LogicalRow(
                identity_state="valid" if valid else "invalid",
                entity_kind="wfm" if valid else "invalid_row",
                canonical_primary_id=task_no if valid else None,
                canonical_parent_rfc_no=rfc_no if valid else None,
                fields=tuple(logical_fields),
                findings=tuple(finding.logical() for finding in row_findings),
            )
            observation = NormalizedObservationEvidence(
                entity_kind=logical_row.entity_kind,
                identity_state=logical_row.identity_state,
                canonical_primary_id=logical_row.canonical_primary_id,
                canonical_parent_rfc_no=logical_row.canonical_parent_rfc_no,
                row_ordinal=row_ordinal,
                sheet_ordinal=matrix.sheet_ordinal,
                row_logical_sha256=row_logical_sha256(logical_row),
                source_row_chronology_utc=source_chronology_utc,
                fields=tuple(_logical_to_normalized(field) for field in logical_fields),
            )
            rows.append(
                ParsedWfmRow(
                    observation=observation,
                    findings=tuple(row_findings),
                    logical_row=logical_row,
                    classification="UNIQUE",
                    duplicate_count=1,
                    conflict_variant_count=1,
                )
            )

        return WfmParseResult(
            source_family=_SOURCE_FAMILY,
            source_profile_id=versions.source_profile_id,
            header_registry_id=versions.header_registry_id,
            vocabulary_registry_id=versions.vocabulary_registry_id,
            parser_profile_id=versions.parser_profile_id,
            source_chronology_utc=source_chronology_utc,
            rows=tuple(_classify_rows(rows)),
            global_findings=global_findings,
        )
    finally:
        workbook.close()


__all__ = ["ParsedWfmRow", "WfmParseResult", "parse_wfm_service_provider"]
