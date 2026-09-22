from __future__ import annotations

import re
from dataclasses import dataclass, replace
from datetime import UTC, date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import unicodedata2
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.cell.read_only import ReadOnlyCell

from soma.foundation.errors import SomaError
from soma.reference.domain.matching import is_match_whitespace, normalize_match_key, trim_match_whitespace
from soma.ticket_import.profiles.registry import ImportFieldSpec, require_field_spec, require_profile_versions
from soma.ticket_import.reconciliation.engine import (
    LogicalField,
    LogicalFinding,
    LogicalRow,
    field_logical_sha256,
    partition_row_variants,
    row_logical_sha256,
)
from soma.ticket_import.repositories.observations import (
    NormalizedFieldEvidence,
    NormalizedObservationEvidence,
)

from .xlsx_security import XlsxPreflightResult

_SOURCE_FAMILY = "advanced_search_sr"
_MAX_WORKSHEETS = 32
_MAX_HEADER_SEARCH_NONEMPTY_ROWS = 100
_MAX_MATRIX_ROWS = 250_000
_MAX_PHYSICAL_COLUMNS = 512
_MAX_LOGICAL_CELLS_TOTAL = 8_000_000
_MAX_SHARED_STRINGS = 1_000_000
_MAX_DEFINED_NAMES = 4_096
_MAX_CELL_UTF8_BYTES = 65_536
_MAX_HEADER_UTF8_BYTES = 1_024
_MAX_FINDING_MESSAGE_UTF8_BYTES = 2_048
_INT64_MAX = 9_223_372_036_854_775_807

_SR_NO = re.compile(r"[0-9]{8}\Z")
_TIMESTAMP = re.compile(
    r"(?P<date>[0-9]{4}-[0-9]{2}-[0-9]{2})[T ]"
    r"(?P<time>[0-9]{2}:[0-9]{2}:[0-9]{2})"
    r"(?:\.(?P<fraction>[0-9]{1,6}))?"
    r"(?P<offset>Z|[+-][0-9]{2}:[0-9]{2})?\Z"
)
_DECIMAL_DAYS = re.compile(r"\+?(?:[0-9]+(?:\.[0-9]+)?|\.[0-9]+)\Z")

_HEADER_ALIASES = {
    "sr_no": ("SRNo",),
    "problem_summary": ("Problem Summary",),
    "report_date": ("Report Date",),
    "customer_contact_label": ("Customer Contact",),
    "customer_severity": ("Customer Severity",),
    "current_handler_label": ("Current Handler",),
    "status": ("Status",),
    "customer_org_label": ("Customer Org.",),
    "customer_account_code": ("Customer Account Code",),
    "suspend_planned_end": ("Suspend Planned End Date",),
    "suspension_duration": ("Suspension Duration",),
    "last_update": ("Last Update",),
    "resolve_by": ("ResolveBy",),
    "resolve_by_suspend": ("Resolve By Suspend",),
    "owner": ("Owner",),
    "l2_assignee": ("L2 Assignee",),
    "l3_assignee": ("L3 Assignee",),
    "related_sr": ("Related SR",),
    "authorization_no": ("Authorization No.",),
    "escalation_no": ("Escalation No",),
    "pr_number": ("Pr Number",),
}
_DISCARDED_ALIASES = ("Product",)

_STATUS_VALUES = (
    ("Closed", "terminal_closed"),
    ("Resolved", "terminal_resolved"),
    ("Cancelled", "terminal_cancelled"),
    ("Customer Agreed Suspend", "active_suspension_status"),
)
_SEVERITY_VALUES = (
    ("Critical", "critical"),
    ("Major", "major"),
    ("Minor", "minor"),
    ("Non-fault inquiry", "non_fault_inquiry"),
)


@dataclass(frozen=True, slots=True)
class ParsedFinding:
    finding_code: str
    severity: str
    scope_kind: str
    message_text: str
    field_key: str | None = None
    sheet_ordinal: int | None = None
    row_ordinal: int | None = None

    def logical(self) -> LogicalFinding:
        return LogicalFinding(
            finding_code=self.finding_code,
            severity=self.severity,
            scope_kind=self.scope_kind,
            field_key=self.field_key,
        )


@dataclass(frozen=True, slots=True)
class ParsedAdvancedSearchRow:
    observation: NormalizedObservationEvidence
    findings: tuple[ParsedFinding, ...]
    logical_row: LogicalRow
    classification: str
    duplicate_count: int
    conflict_variant_count: int


@dataclass(frozen=True, slots=True)
class AdvancedSearchParseResult:
    source_family: str
    source_profile_id: str
    header_registry_id: str
    vocabulary_registry_id: str
    parser_profile_id: str
    rows: tuple[ParsedAdvancedSearchRow, ...]
    global_findings: tuple[ParsedFinding, ...]

    @property
    def observations(self) -> tuple[NormalizedObservationEvidence, ...]:
        return tuple(row.observation for row in self.rows)

    @property
    def logical_rows(self) -> tuple[LogicalRow, ...]:
        return tuple(row.logical_row for row in self.rows)

    @property
    def logical_global_findings(self) -> tuple[LogicalFinding, ...]:
        return tuple(finding.logical() for finding in self.global_findings)


@dataclass(frozen=True, slots=True)
class _Matrix:
    sheet_ordinal: int
    sheet_name: str
    header_row_ordinal: int
    columns: dict[str, int]


def _source_error(code: str, message: str) -> SomaError:
    return SomaError(code, message)


def _bounded_message(message: str) -> str:
    encoded = message.encode("utf-8", errors="strict")
    if len(encoded) <= _MAX_FINDING_MESSAGE_UTF8_BYTES:
        return message
    return encoded[:_MAX_FINDING_MESSAGE_UTF8_BYTES].decode("ascii", errors="strict")


def _finding(
    code: str,
    severity: str,
    scope: str,
    message: str,
    *,
    field_key: str | None = None,
    sheet_ordinal: int | None = None,
    row_ordinal: int | None = None,
) -> ParsedFinding:
    return ParsedFinding(
        code,
        severity,
        scope,
        _bounded_message(message),
        field_key,
        sheet_ordinal,
        row_ordinal,
    )


def _collapse_unicode_whitespace(value: str) -> str:
    normalized = unicodedata2.normalize("NFKC", value)
    output: list[str] = []
    pending_space = False
    for character in normalized:
        if is_match_whitespace(character):
            if output:
                pending_space = True
            continue
        if pending_space:
            output.append(" ")
            pending_space = False
        output.append(character)
    return "".join(output)


def _header_key(raw: object) -> str | None:
    if raw is None:
        return None
    if not isinstance(raw, str):
        raw = _scalar_source_text(raw)
        if raw is None:
            return None
    if "\x00" in raw:
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "workbook header contains NUL")
    if len(raw.encode("utf-8", errors="strict")) > _MAX_HEADER_UTF8_BYTES:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook header exceeds the UTF-8 byte ceiling")
    collapsed = _collapse_unicode_whitespace(raw)
    if not collapsed:
        return None
    return normalize_match_key(
        collapsed,
        raw_max_utf8_bytes=_MAX_HEADER_UTF8_BYTES,
        key_max_utf8_bytes=_MAX_HEADER_UTF8_BYTES,
    )


def _build_alias_registry() -> tuple[dict[str, str], set[str]]:
    aliases: dict[str, str] = {}
    for field_key, values in _HEADER_ALIASES.items():
        for alias in values:
            normalized = _header_key(alias)
            assert normalized is not None
            previous = aliases.setdefault(normalized, field_key)
            if previous != field_key:
                raise RuntimeError("Advanced Search header registry aliases collide")
    discarded: set[str] = set()
    for alias in _DISCARDED_ALIASES:
        normalized = _header_key(alias)
        assert normalized is not None
        discarded.add(normalized)
    return aliases, discarded


_HEADER_KEYS, _DISCARDED_HEADER_KEYS = _build_alias_registry()


def _guayaquil_timezone() -> timezone | ZoneInfo:
    try:
        return ZoneInfo("America/Guayaquil")
    except ZoneInfoNotFoundError:
        return timezone(timedelta(hours=-5), name="America/Guayaquil")


def _is_formula(cell: Cell | ReadOnlyCell) -> bool:
    return getattr(cell, "data_type", None) == "f"


def _scalar_source_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if type(value) is bool:
        return "TRUE" if value else "FALSE"
    if type(value) is int:
        return str(value)
    if type(value) is float:
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return None


def _validate_source_text(source_text: str | None, spec: ImportFieldSpec) -> None:
    if source_text is None:
        return
    if "\x00" in source_text:
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", f"{spec.field_key} contains NUL")
    encoded = source_text.encode("utf-8", errors="strict")
    if len(encoded) > min(spec.max_utf8_bytes, _MAX_CELL_UTF8_BYTES):
        raise _source_error("XLSX_RESOURCE_LIMIT", f"{spec.field_key} exceeds its UTF-8 byte ceiling")
    if spec.max_lines == 1 and ("\r" in source_text or "\n" in source_text):
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", f"{spec.field_key} must be one line")
    if source_text.count("\n") + 1 > spec.max_lines:
        raise _source_error("XLSX_RESOURCE_LIMIT", f"{spec.field_key} exceeds its line ceiling")


def _trim_spreadsheet_whitespace(value: str) -> str:
    return trim_match_whitespace(value)


def _logical_to_normalized(field: LogicalField) -> NormalizedFieldEvidence:
    return NormalizedFieldEvidence(
        field_key=field.field_key,
        field_class=field.field_class,
        value_state=field.value_state,
        value_kind=field.value_kind,
        source_text=field.source_text,
        normalized_text=field.normalized_text,
        integer_value=field.integer_value,
        vocabulary_id=field.vocabulary_id,
        field_logical_sha256=field_logical_sha256(field),
    )


def _blank_field(spec: ImportFieldSpec, *, source_text: str | None = None) -> LogicalField:
    return LogicalField(
        field_key=spec.field_key,
        field_class=spec.field_class,
        value_state="blank",
        value_kind=spec.value_kind,
        vocabulary_id=spec.vocabulary_id,
        source_text=source_text,
    )


def _malformed_field(spec: ImportFieldSpec, source_text: str | None) -> LogicalField:
    return LogicalField(
        field_key=spec.field_key,
        field_class=spec.field_class,
        value_state="malformed",
        value_kind=spec.value_kind,
        vocabulary_id=spec.vocabulary_id,
        source_text=source_text,
    )


def _parse_text_field(spec: ImportFieldSpec, value: object) -> LogicalField:
    source_text = _scalar_source_text(value)
    _validate_source_text(source_text, spec)
    if source_text is None:
        return _blank_field(spec)
    if isinstance(value, (datetime, date, time)) or type(value) is bool:
        return _malformed_field(spec, source_text)
    normalized = _trim_spreadsheet_whitespace(source_text)
    if not normalized:
        return _blank_field(spec, source_text=source_text)
    return LogicalField(
        field_key=spec.field_key,
        field_class=spec.field_class,
        value_state="usable",
        value_kind="text",
        source_text=source_text,
        normalized_text=normalized,
    )


def _controlled_key(source_text: str) -> str:
    collapsed = _collapse_unicode_whitespace(source_text)
    if not collapsed:
        return ""
    return normalize_match_key(
        collapsed,
        raw_max_utf8_bytes=_MAX_CELL_UTF8_BYTES,
        key_max_utf8_bytes=_MAX_CELL_UTF8_BYTES,
    )


def _controlled_registry(field_key: str) -> dict[str, str]:
    if field_key == "status":
        values = _STATUS_VALUES
    elif field_key == "customer_severity":
        values = _SEVERITY_VALUES
    else:
        raise RuntimeError("unknown Advanced Search controlled field")
    output: dict[str, str] = {}
    for canonical, _resolved in values:
        output[_controlled_key(canonical)] = canonical
    return output


_STATUS_REGISTRY = _controlled_registry("status")
_SEVERITY_REGISTRY = _controlled_registry("customer_severity")


def _parse_controlled_field(
    spec: ImportFieldSpec,
    value: object,
    *,
    sheet_ordinal: int,
    row_ordinal: int,
) -> tuple[LogicalField, tuple[ParsedFinding, ...]]:
    source_text = _scalar_source_text(value)
    _validate_source_text(source_text, spec)
    if source_text is None:
        return _blank_field(spec), ()
    if isinstance(value, (datetime, date, time)) or type(value) is bool:
        return _malformed_field(spec, source_text), ()
    trimmed = _trim_spreadsheet_whitespace(source_text)
    if not trimmed:
        return _blank_field(spec, source_text=source_text), ()
    key = _controlled_key(trimmed)
    registry = _STATUS_REGISTRY if spec.field_key == "status" else _SEVERITY_REGISTRY
    canonical = registry.get(key)
    if canonical is None:
        field = LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="unknown",
            value_kind="controlled",
            vocabulary_id=spec.vocabulary_id,
            source_text=source_text,
        )
        finding = _finding(
            "SOURCE_CONTROLLED_VALUE_UNKNOWN",
            "warning",
            "field",
            f"{spec.field_key} is outside the accepted controlled vocabulary",
            field_key=spec.field_key,
            sheet_ordinal=sheet_ordinal,
            row_ordinal=row_ordinal,
        )
        return field, (finding,)
    return (
        LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="usable",
            value_kind="controlled",
            vocabulary_id=spec.vocabulary_id,
            source_text=source_text,
            normalized_text=canonical,
        ),
        (),
    )


def _parse_offset(token: str | None) -> timezone | ZoneInfo:
    if token is None:
        return _guayaquil_timezone()
    if token == "Z":
        return UTC
    sign = 1 if token[0] == "+" else -1
    hours = int(token[1:3])
    minutes = int(token[4:6])
    if hours > 23 or minutes > 59:
        raise ValueError("invalid offset")
    total = sign * (hours * 60 + minutes)
    if abs(total) >= 24 * 60:
        raise ValueError("invalid offset")
    return timezone(timedelta(minutes=total))


def _datetime_to_epoch_seconds(value: datetime) -> int:
    aware = value
    if aware.tzinfo is None:
        aware = aware.replace(tzinfo=_guayaquil_timezone())
    utc_value = aware.astimezone(UTC).replace(microsecond=0)
    try:
        epoch = int(utc_value.timestamp())
    except (OverflowError, OSError, ValueError) as exc:
        raise ValueError("unsupported timestamp") from exc
    if epoch < 0:
        raise ValueError("pre-epoch timestamp")
    return epoch


def _parse_timestamp_text(source_text: str) -> int:
    trimmed = _trim_spreadsheet_whitespace(source_text)
    match = _TIMESTAMP.fullmatch(trimmed)
    if match is None:
        raise ValueError("timestamp grammar mismatch")
    fraction = match.group("fraction") or ""
    microsecond = int(fraction.ljust(6, "0")) if fraction else 0
    try:
        base = datetime.strptime(
            f"{match.group('date')} {match.group('time')}",
            "%Y-%m-%d %H:%M:%S",
        ).replace(microsecond=microsecond)
        zone = _parse_offset(match.group("offset"))
        aware = base.replace(tzinfo=zone, fold=0)
        if isinstance(zone, ZoneInfo):
            alternate = base.replace(tzinfo=zone, fold=1)
            if aware.utcoffset() != alternate.utcoffset():
                raise ValueError("ambiguous local timestamp")
            if aware.astimezone(UTC).astimezone(zone).replace(tzinfo=None) != base:
                raise ValueError("nonexistent local timestamp")
        return _datetime_to_epoch_seconds(aware)
    except (OverflowError, ValueError) as exc:
        raise ValueError("invalid timestamp") from exc


def _parse_instant_field(spec: ImportFieldSpec, value: object) -> LogicalField:
    source_text = _scalar_source_text(value)
    _validate_source_text(source_text, spec)
    if value is None:
        return _blank_field(spec)
    try:
        if isinstance(value, datetime):
            epoch = _datetime_to_epoch_seconds(value)
        elif isinstance(value, date) and not isinstance(value, datetime):
            epoch = _datetime_to_epoch_seconds(datetime.combine(value, time.min))
        elif isinstance(value, str):
            if not _trim_spreadsheet_whitespace(value):
                return _blank_field(spec, source_text=value)
            epoch = _parse_timestamp_text(value)
        else:
            return _malformed_field(spec, source_text)
    except ValueError:
        return _malformed_field(spec, source_text)
    assert source_text is not None
    return LogicalField(
        field_key=spec.field_key,
        field_class=spec.field_class,
        value_state="usable",
        value_kind="instant",
        source_text=source_text,
        integer_value=epoch,
    )


def _duration_decimal(value: object) -> tuple[Decimal | None, str | None]:
    source_text = _scalar_source_text(value)
    if value is None:
        return None, source_text
    if type(value) is bool or isinstance(value, (datetime, date, time)):
        raise ValueError("unsupported duration type")
    if isinstance(value, str):
        trimmed = _trim_spreadsheet_whitespace(value)
        if not trimmed:
            return None, value
        if _DECIMAL_DAYS.fullmatch(trimmed) is None:
            raise ValueError("duration grammar mismatch")
        return Decimal(trimmed), value
    if type(value) in {int, float}:
        token = str(value)
        return Decimal(token), token
    raise ValueError("unsupported duration type")


def _parse_duration_field(spec: ImportFieldSpec, value: object) -> LogicalField:
    source_text = _scalar_source_text(value)
    _validate_source_text(source_text, spec)
    if value is None or isinstance(value, str) and not _trim_spreadsheet_whitespace(value):
        return LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="usable",
            value_kind="duration_seconds",
            source_text="" if source_text is None else source_text,
            integer_value=0,
        )
    try:
        decimal_days, source_evidence = _duration_decimal(value)
        if decimal_days is None or not decimal_days.is_finite() or decimal_days < 0:
            raise ValueError("duration invalid")
        seconds = (decimal_days * Decimal(86_400)).quantize(Decimal(1), rounding=ROUND_HALF_EVEN)
        integer_value = int(seconds)
        if integer_value < 0 or integer_value > _INT64_MAX:
            raise ValueError("duration out of range")
        assert source_evidence is not None
        return LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="usable",
            value_kind="duration_seconds",
            source_text=source_evidence,
            integer_value=integer_value,
        )
    except (InvalidOperation, OverflowError, ValueError):
        return _malformed_field(spec, source_text)


def _cell_is_physically_nonempty(cell: Cell | ReadOnlyCell) -> bool:
    return cell.value is not None


def _row_is_physically_nonempty(row: Iterable[Cell | ReadOnlyCell]) -> bool:
    return any(_cell_is_physically_nonempty(cell) for cell in row)


def _resolve_header_row(
    row: tuple[Cell | ReadOnlyCell, ...],
    *,
    sheet_ordinal: int,
    row_ordinal: int,
) -> dict[str, int] | None:
    resolved: dict[str, int] = {}
    for column_ordinal, cell in enumerate(row, start=1):
        if column_ordinal > _MAX_PHYSICAL_COLUMNS:
            raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
        key = _header_key(cell.value)
        if key is None or key in _DISCARDED_HEADER_KEYS:
            continue
        field_key = _HEADER_KEYS.get(key)
        if field_key is None:
            continue
        if field_key in resolved:
            raise _source_error(
                "SOURCE_DUPLICATE_SEMANTIC_HEADER",
                f"worksheet {sheet_ordinal} row {row_ordinal} repeats semantic header {field_key}",
            )
        resolved[field_key] = column_ordinal
    return resolved if "sr_no" in resolved else None


def _discover_matrix(workbook) -> _Matrix:
    if len(workbook.worksheets) > _MAX_WORKSHEETS:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the worksheet-count ceiling")
    if len(getattr(workbook, "shared_strings", ())) > _MAX_SHARED_STRINGS:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the shared-string ceiling")
    if len(getattr(workbook, "defined_names", ())) > _MAX_DEFINED_NAMES:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the defined-name ceiling")

    candidates: list[_Matrix] = []
    scanned_cells = 0
    for sheet_ordinal, worksheet in enumerate(workbook.worksheets, start=1):
        nonempty_seen = 0
        for row in worksheet.iter_rows():
            if len(row) > _MAX_PHYSICAL_COLUMNS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                scanned_cells += 1
                if scanned_cells > _MAX_LOGICAL_CELLS_TOTAL:
                    raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the logical-cell ceiling")
                if isinstance(value, str):
                    if "\x00" in value:
                        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "workbook cell text contains NUL")
                    if len(value.encode("utf-8", errors="strict")) > _MAX_CELL_UTF8_BYTES:
                        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook cell exceeds the UTF-8 byte ceiling")
            if not _row_is_physically_nonempty(row):
                continue
            nonempty_seen += 1
            if nonempty_seen > _MAX_HEADER_SEARCH_NONEMPTY_ROWS:
                continue
            row_ordinal = int(row[0].row) if row else nonempty_seen
            resolved = _resolve_header_row(
                row,
                sheet_ordinal=sheet_ordinal,
                row_ordinal=row_ordinal,
            )
            if resolved is not None:
                candidates.append(
                    _Matrix(
                        sheet_ordinal=sheet_ordinal,
                        sheet_name=worksheet.title,
                        header_row_ordinal=row_ordinal,
                        columns=resolved,
                    )
                )
    if not candidates:
        raise _source_error("SOURCE_MATRIX_NOT_FOUND", "no Advanced Search matrix with SRNo was found")
    if len(candidates) != 1:
        raise _source_error("SOURCE_MATRIX_AMBIGUOUS", "multiple Advanced Search matrices contain SRNo")
    matrix = candidates[0]
    if "sr_no" not in matrix.columns:
        raise _source_error("SOURCE_REQUIRED_HEADER_MISSING", "Advanced Search SRNo header is missing")
    return matrix


def _global_coverage_findings(matrix: _Matrix) -> tuple[ParsedFinding, ...]:
    findings: list[ParsedFinding] = []
    for field_key in sorted(_HEADER_ALIASES):
        if field_key == "sr_no" or field_key in matrix.columns:
            continue
        findings.append(
            _finding(
                "SOURCE_HEADER_COVERAGE_MISSING",
                "warning",
                "workbook",
                f"registered Advanced Search header {field_key} is absent",
                field_key=field_key,
            )
        )
    return tuple(findings)


def _parse_identity(value: object) -> str | None:
    if isinstance(value, str):
        candidate = _trim_spreadsheet_whitespace(value)
    elif type(value) is int:
        candidate = str(value)
    else:
        return None
    return candidate if _SR_NO.fullmatch(candidate) is not None else None


def _parse_semantic_field(
    field_key: str,
    cell: Cell | ReadOnlyCell,
    *,
    sheet_ordinal: int,
    row_ordinal: int,
) -> tuple[LogicalField, tuple[ParsedFinding, ...]]:
    spec = require_field_spec(_SOURCE_FAMILY, field_key)
    if _is_formula(cell):
        raise _source_error(
            "SOURCE_FORMULA_IN_SEMANTIC_FIELD",
            f"registered semantic field {field_key} contains a formula",
        )
    if spec.value_kind == "text":
        return _parse_text_field(spec, cell.value), ()
    if spec.value_kind == "controlled":
        return _parse_controlled_field(
            spec,
            cell.value,
            sheet_ordinal=sheet_ordinal,
            row_ordinal=row_ordinal,
        )
    if spec.value_kind == "instant":
        return _parse_instant_field(spec, cell.value), ()
    if spec.value_kind == "duration_seconds":
        return _parse_duration_field(spec, cell.value), ()
    raise RuntimeError("unsupported Advanced Search field kind")


def _classify_rows(rows: list[ParsedAdvancedSearchRow]) -> list[ParsedAdvancedSearchRow]:
    valid_rows = [row for row in rows if row.logical_row.identity_state == "valid"]
    variants = partition_row_variants(row.logical_row for row in valid_rows)
    classification_by_key: dict[tuple[str, str | None, str], tuple[str, int, int]] = {}
    for variant in variants:
        classification_by_key[
            (
                variant.row.entity_kind,
                variant.row.canonical_primary_id,
                variant.row_logical_sha256,
            )
        ] = (
            variant.classification,
            variant.duplicate_count,
            variant.conflict_variant_count,
        )
    output: list[ParsedAdvancedSearchRow] = []
    for parsed in rows:
        if parsed.logical_row.identity_state != "valid":
            output.append(parsed)
            continue
        key = (
            parsed.logical_row.entity_kind,
            parsed.logical_row.canonical_primary_id,
            parsed.observation.row_logical_sha256,
        )
        classification, duplicate_count, conflict_count = classification_by_key[key]
        findings = list(parsed.findings)
        if classification == "EQUIVALENT_DUPLICATE":
            findings.append(
                _finding(
                    "SOURCE_EQUIVALENT_DUPLICATE",
                    "warning",
                    "row",
                    "equivalent Advanced Search rows share one logical proposal variant",
                    sheet_ordinal=parsed.observation.sheet_ordinal,
                    row_ordinal=parsed.observation.row_ordinal,
                )
            )
        elif classification == "CONFLICT_MEMBER":
            findings.append(
                _finding(
                    "SOURCE_IDENTITY_CONFLICT",
                    "error",
                    "identity",
                    "same Service Request identity has conflicting logical source rows",
                    sheet_ordinal=parsed.observation.sheet_ordinal,
                    row_ordinal=parsed.observation.row_ordinal,
                )
            )
        output.append(
            replace(
                parsed,
                findings=tuple(findings),
                classification=classification,
                duplicate_count=duplicate_count,
                conflict_variant_count=conflict_count,
            )
        )
    return output


def parse_advanced_search(
    path: Path,
    *,
    preflight: XlsxPreflightResult,
) -> AdvancedSearchParseResult:
    requested = Path(path)
    try:
        requested_resolved = requested.resolve(strict=True)
        preflight_resolved = Path(preflight.path).resolve(strict=True)
    except OSError as exc:
        preflight.close()
        raise _source_error(
            "IMPORT_SOURCE_UNAVAILABLE",
            "preflighted workbook is no longer available",
        ) from exc
    if requested_resolved != preflight_resolved:
        preflight.close()
        raise _source_error(
            "IMPORT_SOURCE_PROFILE_MISMATCH",
            "parser path does not match the preflighted workbook identity",
        )

    versions = require_profile_versions(_SOURCE_FAMILY)
    try:
        workbook = load_workbook(
            preflight.semantic_stream(),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        preflight.close()
        raise _source_error(
            "IMPORT_SOURCE_PROFILE_MISMATCH",
            "preflighted workbook could not be opened semantically",
        ) from exc

    try:
        matrix = _discover_matrix(workbook)
        global_findings = _global_coverage_findings(matrix)
        worksheet = workbook.worksheets[matrix.sheet_ordinal - 1]
        rows: list[ParsedAdvancedSearchRow] = []
        logical_cells = 0
        matrix_rows = 0
        max_semantic_column = max(matrix.columns.values())

        for physical in worksheet.iter_rows(min_row=matrix.header_row_ordinal + 1):
            if len(physical) > _MAX_PHYSICAL_COLUMNS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
            if not _row_is_physically_nonempty(physical):
                continue
            matrix_rows += 1
            if matrix_rows > _MAX_MATRIX_ROWS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "Advanced Search matrix exceeds the row ceiling")
            row_ordinal = int(physical[0].row) if physical else matrix.header_row_ordinal + matrix_rows
            semantic_count = len(matrix.columns)
            logical_cells += semantic_count
            if logical_cells > _MAX_LOGICAL_CELLS_TOTAL:
                raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the logical-cell ceiling")

            if len(physical) < max_semantic_column:
                cells_by_column = {index + 1: cell for index, cell in enumerate(physical)}
            else:
                cells_by_column = {index + 1: cell for index, cell in enumerate(physical)}

            identity_column = matrix.columns["sr_no"]
            identity_cell = cells_by_column.get(identity_column)
            if identity_cell is not None and _is_formula(identity_cell):
                raise _source_error(
                    "SOURCE_FORMULA_IN_SEMANTIC_FIELD",
                    "registered semantic field sr_no contains a formula",
                )
            identity = _parse_identity(None if identity_cell is None else identity_cell.value)
            identity_valid = identity is not None
            row_findings: list[ParsedFinding] = []
            if not identity_valid:
                row_findings.append(
                    _finding(
                        "SR_ID_INVALID",
                        "error",
                        "identity",
                        "Advanced Search SRNo is not exactly eight ASCII digits",
                        field_key="sr_no",
                        sheet_ordinal=matrix.sheet_ordinal,
                        row_ordinal=row_ordinal,
                    )
                )

            logical_fields: list[LogicalField] = []
            for field_key, column_ordinal in matrix.columns.items():
                if field_key == "sr_no":
                    continue
                cell = cells_by_column.get(column_ordinal)
                if cell is None:
                    class _Blank:
                        value = None
                        data_type = "n"
                    cell = _Blank()  # type: ignore[assignment]
                field, findings = _parse_semantic_field(
                    field_key,
                    cell,  # type: ignore[arg-type]
                    sheet_ordinal=matrix.sheet_ordinal,
                    row_ordinal=row_ordinal,
                )
                logical_fields.append(field)
                row_findings.extend(findings)

            logical_row = LogicalRow(
                identity_state="valid" if identity_valid else "invalid",
                entity_kind="service_request" if identity_valid else "invalid_row",
                canonical_primary_id=identity,
                canonical_parent_rfc_no=None,
                fields=tuple(logical_fields),
                findings=tuple(finding.logical() for finding in row_findings),
            )
            row_hash = row_logical_sha256(logical_row)
            chronology: int | None = None
            for field in logical_fields:
                if field.field_key == "last_update" and field.value_state == "usable":
                    chronology = field.integer_value
                    break
            observation = NormalizedObservationEvidence(
                entity_kind=logical_row.entity_kind,
                identity_state=logical_row.identity_state,
                canonical_primary_id=logical_row.canonical_primary_id,
                canonical_parent_rfc_no=None,
                row_ordinal=row_ordinal,
                sheet_ordinal=matrix.sheet_ordinal,
                row_logical_sha256=row_hash,
                source_row_chronology_utc=chronology,
                fields=tuple(_logical_to_normalized(field) for field in logical_fields),
            )
            rows.append(
                ParsedAdvancedSearchRow(
                    observation=observation,
                    findings=tuple(row_findings),
                    logical_row=logical_row,
                    classification="UNIQUE",
                    duplicate_count=1,
                    conflict_variant_count=1,
                )
            )

        classified = _classify_rows(rows)
        return AdvancedSearchParseResult(
            source_family=_SOURCE_FAMILY,
            source_profile_id=versions.source_profile_id,
            header_registry_id=versions.header_registry_id,
            vocabulary_registry_id=versions.vocabulary_registry_id,
            parser_profile_id=versions.parser_profile_id,
            rows=tuple(classified),
            global_findings=global_findings,
        )
    finally:
        try:
            workbook.close()
        finally:
            preflight.close()