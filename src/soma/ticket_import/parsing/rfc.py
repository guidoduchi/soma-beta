from __future__ import annotations

import re
from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import load_workbook

from soma.foundation.errors import SomaError
from soma.ticket_import.profiles.registry import require_field_spec, require_profile_versions
from soma.ticket_import.reconciliation.engine import LogicalField, LogicalRow, partition_row_variants, row_logical_sha256
from soma.ticket_import.repositories.observations import NormalizedObservationEvidence
from soma.tickets.validation import validate_rfc_no

from .advanced_search import (
    ParsedFinding,
    _MAX_CELL_UTF8_BYTES,
    _MAX_DEFINED_NAMES,
    _MAX_HEADER_SEARCH_NONEMPTY_ROWS,
    _MAX_LOGICAL_CELLS_TOTAL,
    _MAX_MATRIX_ROWS,
    _MAX_PHYSICAL_COLUMNS,
    _MAX_SHARED_STRINGS,
    _MAX_WORKSHEETS,
    _Matrix,
    _blank_field,
    _controlled_key,
    _finding,
    _header_key,
    _is_formula,
    _logical_to_normalized,
    _malformed_field,
    _parse_instant_field,
    _parse_text_field,
    _row_is_physically_nonempty,
    _scalar_source_text,
    _source_error,
    _trim_spreadsheet_whitespace,
    _validate_source_text,
)
from .xlsx_security import XlsxPreflightResult

_SOURCE_FAMILY = "rfc_enhanced"
_RFC_PREFIX = re.compile(r"NC[0-9]{14}")
_HEADER_ALIASES = {
    "rfc_no": ("Task ID",),
    "external_created_at": ("Create Time",),
    "creator": ("Creator",),
    "customer_account_number": ("Customer Account Number",),
    "customer_account_name": ("Customer Account Name",),
    "severity": ("Severity",),
    "summary": ("Summary",),
    "status": ("Status",),
    "owner_external_id": ("Owner",),
    "owner_name": ("Owner Name",),
    "l1_handler_name": ("L1 Handler Name",),
    "l2_handler_name": ("L2 Handler Name",),
    "last_update": ("Last Update Time",),
    "service_type": ("Service Type",),
    "scenario": ("Scenario",),
    "product_line_code": ("Product Line Code",),
    "product_line": ("Product Line",),
    "product_code": ("Product Code",),
    "product_name": ("Product Name",),
}
_STATUS_VALUES = ("Implement", "Closed", "Cancelled")


@dataclass(frozen=True, slots=True)
class ParsedRfcRow:
    observation: NormalizedObservationEvidence
    findings: tuple[ParsedFinding, ...]
    logical_row: LogicalRow
    classification: str
    duplicate_count: int
    conflict_variant_count: int


@dataclass(frozen=True, slots=True)
class RfcParseResult:
    source_family: str
    source_profile_id: str
    header_registry_id: str
    vocabulary_registry_id: str
    parser_profile_id: str
    rows: tuple[ParsedRfcRow, ...]
    global_findings: tuple[ParsedFinding, ...]

    @property
    def observations(self) -> tuple[NormalizedObservationEvidence, ...]:
        return tuple(row.observation for row in self.rows)


@dataclass(frozen=True, slots=True)
class _RfcIdentity:
    state: str
    canonical: str | None
    finding_code: str | None


def _build_header_registry() -> dict[str, str]:
    output: dict[str, str] = {}
    for field_key, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            key = _header_key(alias)
            assert key is not None
            previous = output.setdefault(key, field_key)
            if previous != field_key:
                raise RuntimeError("RFC header registry aliases collide")
    return output


_HEADER_KEYS = _build_header_registry()
_STATUS_REGISTRY = {_controlled_key(value): value for value in _STATUS_VALUES}


def _resolve_header_row(row, *, sheet_ordinal: int, row_ordinal: int) -> dict[str, int] | None:
    resolved: dict[str, int] = {}
    for column_ordinal, cell in enumerate(row, start=1):
        if column_ordinal > _MAX_PHYSICAL_COLUMNS:
            raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
        key = _header_key(cell.value)
        if key is None:
            continue
        field_key = _HEADER_KEYS.get(key)
        if field_key is None:
            continue
        if field_key in resolved:
            raise _source_error(
                "SOURCE_DUPLICATE_SEMANTIC_HEADER",
                f"worksheet {sheet_ordinal} row {row_ordinal} repeats semantic header {field_key}",
            )
        resolved[field_key] = column_ordinal
    return resolved if "rfc_no" in resolved else None


def _discover_matrix(workbook) -> _Matrix:
    if len(workbook.worksheets) > _MAX_WORKSHEETS:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the worksheet-count ceiling")
    if len(getattr(workbook, "shared_strings", ())) > _MAX_SHARED_STRINGS:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the shared-string ceiling")
    if len(getattr(workbook, "defined_names", ())) > _MAX_DEFINED_NAMES:
        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the defined-name ceiling")

    candidates: list[_Matrix] = []
    scanned_cells = 0
    for sheet_ordinal, worksheet in enumerate(workbook.worksheets, start=1):
        nonempty_seen = 0
        for row in worksheet.iter_rows():
            if len(row) > _MAX_PHYSICAL_COLUMNS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                scanned_cells += 1
                if scanned_cells > _MAX_LOGICAL_CELLS_TOTAL:
                    raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the logical-cell ceiling")
                if isinstance(value, str):
                    if "\x00" in value:
                        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "workbook cell text contains NUL")
                    if len(value.encode("utf-8", errors="strict")) > _MAX_CELL_UTF8_BYTES:
                        raise _source_error("XLSX_RESOURCE_LIMIT", "workbook cell exceeds the UTF-8 byte ceiling")
            if not _row_is_physically_nonempty(row):
                continue
            nonempty_seen += 1
            if nonempty_seen > _MAX_HEADER_SEARCH_NONEMPTY_ROWS:
                continue
            row_ordinal = int(row[0].row) if row else nonempty_seen
            columns = _resolve_header_row(row, sheet_ordinal=sheet_ordinal, row_ordinal=row_ordinal)
            if columns is not None:
                candidates.append(_Matrix(sheet_ordinal, worksheet.title, row_ordinal, columns))
    if not candidates:
        raise _source_error("SOURCE_MATRIX_NOT_FOUND", "no Enhanced RFC matrix with Task ID was found")
    if len(candidates) != 1:
        raise _source_error("SOURCE_MATRIX_AMBIGUOUS", "multiple Enhanced RFC matrices contain Task ID")
    return candidates[0]


def _global_coverage_findings(matrix: _Matrix) -> tuple[ParsedFinding, ...]:
    return tuple(
        _finding(
            "SOURCE_HEADER_COVERAGE_MISSING",
            "warning",
            "workbook",
            f"registered Enhanced RFC header {field_key} is absent",
            field_key=field_key,
        )
        for field_key in sorted(_HEADER_ALIASES)
        if field_key != "rfc_no" and field_key not in matrix.columns
    )


def _parse_identity(value: object) -> _RfcIdentity:
    source_text = _scalar_source_text(value)
    if source_text is None or isinstance(value, bool):
        return _RfcIdentity("invalid", None, "RFC_ID_INVALID")
    candidate = _trim_spreadsheet_whitespace(source_text)
    if not candidate:
        return _RfcIdentity("invalid", None, "RFC_ID_INVALID")
    try:
        return _RfcIdentity("valid", validate_rfc_no(candidate), None)
    except SomaError:
        prefix = _RFC_PREFIX.match(candidate)
        if prefix is not None and prefix.end() == 16:
            suffix = candidate[16:]
            if suffix and any(not character.isspace() for character in suffix):
                return _RfcIdentity("artifact", None, "RFC_SOURCE_BRANCH_ARTIFACT")
        return _RfcIdentity("invalid", None, "RFC_ID_INVALID")


def _parse_controlled_status(spec, value: object, *, sheet_ordinal: int, row_ordinal: int):
    source_text = _scalar_source_text(value)
    _validate_source_text(source_text, spec)
    if source_text is None:
        return _blank_field(spec), ()
    if not isinstance(value, (str, int, float)) or isinstance(value, bool):
        return _malformed_field(spec, source_text), ()
    trimmed = _trim_spreadsheet_whitespace(source_text)
    if not trimmed:
        return _blank_field(spec, source_text=source_text), ()
    canonical = _STATUS_REGISTRY.get(_controlled_key(trimmed))
    if canonical is None:
        field = LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="unknown",
            value_kind="controlled",
            vocabulary_id=spec.vocabulary_id,
            source_text=source_text,
        )
        return field, (
            _finding(
                "SOURCE_CONTROLLED_VALUE_UNKNOWN",
                "warning",
                "field",
                "status is outside RFC_STATUS_V1",
                field_key="status",
                sheet_ordinal=sheet_ordinal,
                row_ordinal=row_ordinal,
            ),
        )
    return (
        LogicalField(
            field_key=spec.field_key,
            field_class=spec.field_class,
            value_state="usable",
            value_kind="controlled",
            vocabulary_id=spec.vocabulary_id,
            source_text=source_text,
            normalized_text=canonical,
        ),
        (),
    )


def _parse_field(field_key: str, cell, *, sheet_ordinal: int, row_ordinal: int):
    spec = require_field_spec(_SOURCE_FAMILY, field_key)
    if _is_formula(cell):
        raise _source_error(
            "SOURCE_FORMULA_IN_SEMANTIC_FIELD",
            f"registered semantic field {field_key} contains a formula",
        )
    if spec.value_kind == "text":
        return _parse_text_field(spec, cell.value), ()
    if spec.value_kind == "controlled":
        return _parse_controlled_status(spec, cell.value, sheet_ordinal=sheet_ordinal, row_ordinal=row_ordinal)
    if spec.value_kind == "instant":
        return _parse_instant_field(spec, cell.value), ()
    raise RuntimeError("unsupported Enhanced RFC field kind")


def _classify_rows(rows: list[ParsedRfcRow]) -> list[ParsedRfcRow]:
    valid_rows = [row for row in rows if row.logical_row.identity_state == "valid"]
    variants = partition_row_variants(row.logical_row for row in valid_rows)
    by_key = {
        (variant.row.canonical_primary_id, variant.row_logical_sha256): (
            variant.classification,
            variant.duplicate_count,
            variant.conflict_variant_count,
        )
        for variant in variants
    }
    output: list[ParsedRfcRow] = []
    for parsed in rows:
        if parsed.logical_row.identity_state != "valid":
            output.append(parsed)
            continue
        classification, duplicate_count, conflict_count = by_key[
            (parsed.logical_row.canonical_primary_id, parsed.observation.row_logical_sha256)
        ]
        findings = list(parsed.findings)
        if classification == "EQUIVALENT_DUPLICATE":
            findings.append(
                _finding(
                    "SOURCE_EQUIVALENT_DUPLICATE",
                    "warning",
                    "row",
                    "equivalent Enhanced RFC rows share one logical proposal variant",
                    sheet_ordinal=parsed.observation.sheet_ordinal,
                    row_ordinal=parsed.observation.row_ordinal,
                )
            )
        elif classification == "CONFLICT_MEMBER":
            findings.append(
                _finding(
                    "SOURCE_IDENTITY_CONFLICT",
                    "error",
                    "identity",
                    "same RFC identity has conflicting logical source rows",
                    sheet_ordinal=parsed.observation.sheet_ordinal,
                    row_ordinal=parsed.observation.row_ordinal,
                )
            )
        output.append(
            replace(
                parsed,
                findings=tuple(findings),
                classification=classification,
                duplicate_count=duplicate_count,
                conflict_variant_count=conflict_count,
            )
        )
    return output


def parse_rfc_enhanced(path: Path, *, preflight: XlsxPreflightResult) -> RfcParseResult:
    requested = Path(path)
    try:
        requested_resolved = requested.resolve(strict=True)
        preflight_resolved = Path(preflight.path).resolve(strict=True)
    except OSError as exc:
        raise _source_error("IMPORT_SOURCE_UNAVAILABLE", "preflighted workbook is no longer available") from exc
    if requested_resolved != preflight_resolved:
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "parser path does not match preflighted workbook")

    versions = require_profile_versions(_SOURCE_FAMILY)
    try:
        workbook = load_workbook(requested_resolved, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise _source_error("IMPORT_SOURCE_PROFILE_MISMATCH", "preflighted workbook could not be opened semantically") from exc

    try:
        matrix = _discover_matrix(workbook)
        global_findings = _global_coverage_findings(matrix)
        worksheet = workbook.worksheets[matrix.sheet_ordinal - 1]
        rows: list[ParsedRfcRow] = []
        logical_cells = 0
        matrix_rows = 0
        for physical in worksheet.iter_rows(min_row=matrix.header_row_ordinal + 1):
            if len(physical) > _MAX_PHYSICAL_COLUMNS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "worksheet exceeds the physical-column ceiling")
            if not _row_is_physically_nonempty(physical):
                continue
            matrix_rows += 1
            if matrix_rows > _MAX_MATRIX_ROWS:
                raise _source_error("XLSX_RESOURCE_LIMIT", "Enhanced RFC matrix exceeds the row ceiling")
            logical_cells += len(matrix.columns)
            if logical_cells > _MAX_LOGICAL_CELLS_TOTAL:
                raise _source_error("XLSX_RESOURCE_LIMIT", "workbook exceeds the logical-cell ceiling")
            row_ordinal = int(physical[0].row) if physical else matrix.header_row_ordinal + matrix_rows
            cells = {index + 1: cell for index, cell in enumerate(physical)}

            identity_cell = cells.get(matrix.columns["rfc_no"])
            if identity_cell is not None and _is_formula(identity_cell):
                raise _source_error("SOURCE_FORMULA_IN_SEMANTIC_FIELD", "registered semantic field rfc_no contains a formula")
            identity = _parse_identity(None if identity_cell is None else identity_cell.value)
            row_findings: list[ParsedFinding] = []
            if identity.finding_code == "RFC_SOURCE_BRANCH_ARTIFACT":
                row_findings.append(
                    _finding(
                        identity.finding_code,
                        "warning",
                        "identity",
                        "recognized RFC source branch artifact is retained only as skipped evidence",
                        field_key="rfc_no",
                        sheet_ordinal=matrix.sheet_ordinal,
                        row_ordinal=row_ordinal,
                    )
                )
            elif identity.finding_code is not None:
                row_findings.append(
                    _finding(
                        identity.finding_code,
                        "error",
                        "identity",
                        "RFC identity is not exact NC followed by fourteen ASCII digits",
                        field_key="rfc_no",
                        sheet_ordinal=matrix.sheet_ordinal,
                        row_ordinal=row_ordinal,
                    )
                )

            logical_fields: list[LogicalField] = []
            for field_key, column_ordinal in matrix.columns.items():
                if field_key == "rfc_no":
                    continue
                cell = cells.get(column_ordinal)
                if cell is None:
                    class _Blank:
                        value = None
                        data_type = "n"
                    cell = _Blank()
                field, findings = _parse_field(
                    field_key,
                    cell,
                    sheet_ordinal=matrix.sheet_ordinal,
                    row_ordinal=row_ordinal,
                )
                logical_fields.append(field)
                row_findings.extend(findings)

            valid = identity.state == "valid"
            logical_row = LogicalRow(
                identity_state="valid" if valid else "invalid",
                entity_kind="rfc" if valid else "invalid_row",
                canonical_primary_id=identity.canonical,
                canonical_parent_rfc_no=None,
                fields=tuple(logical_fields),
                findings=tuple(finding.logical() for finding in row_findings),
            )
            chronology = next(
                (
                    field.integer_value
                    for field in logical_fields
                    if field.field_key == "last_update" and field.value_state == "usable"
                ),
                None,
            )
            observation = NormalizedObservationEvidence(
                entity_kind=logical_row.entity_kind,
                identity_state=logical_row.identity_state,
                canonical_primary_id=logical_row.canonical_primary_id,
                canonical_parent_rfc_no=None,
                row_ordinal=row_ordinal,
                sheet_ordinal=matrix.sheet_ordinal,
                row_logical_sha256=row_logical_sha256(logical_row),
                source_row_chronology_utc=chronology,
                fields=tuple(_logical_to_normalized(field) for field in logical_fields),
            )
            rows.append(
                ParsedRfcRow(
                    observation=observation,
                    findings=tuple(row_findings),
                    logical_row=logical_row,
                    classification="UNIQUE",
                    duplicate_count=1,
                    conflict_variant_count=1,
                )
            )

        return RfcParseResult(
            source_family=_SOURCE_FAMILY,
            source_profile_id=versions.source_profile_id,
            header_registry_id=versions.header_registry_id,
            vocabulary_registry_id=versions.vocabulary_registry_id,
            parser_profile_id=versions.parser_profile_id,
            rows=tuple(_classify_rows(rows)),
            global_findings=global_findings,
        )
    finally:
        workbook.close()


__all__ = ["ParsedRfcRow", "RfcParseResult", "parse_rfc_enhanced"]
