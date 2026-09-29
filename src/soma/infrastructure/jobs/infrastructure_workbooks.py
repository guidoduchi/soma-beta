from __future__ import annotations

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import os
from pathlib import Path
import posixpath
import tempfile
from typing import BinaryIO
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell

from soma.foundation.contracts.foundation import DurableJobClaim
from soma.foundation.errors import IntegrityFailure, JobClaimConflict, SomaError, ValidationError
from soma.foundation.identifiers import new_uuid4, require_uuid4
from soma.foundation.jobs import DurableJobCoordinator, JobTypeRegistry
from soma.foundation.persistence.connections import ConnectionFactory
from soma.foundation.persistence.uow import ReadSnapshot, UnitOfWork
from soma.foundation.queries.data_instance_identity import DataInstanceIdentityReader
from soma.foundation.strict_json import canonical_json_bytes, loads_canonical_json, loads_strict
from soma.infrastructure.contracts.infrastructure import validate_value
from soma.infrastructure.domain.workbooks import (
    FORMAT_ID, HEADERS, METADATA_KEYS, MODES, SHEET_ORDER, WORKBOOK_VERSION,
    artifact_filename,
)
from soma.infrastructure.jobs import (
    EXPORT_JOB_TYPE, INFRASTRUCTURE_JOB_CONTRACTS,
    validate_export_checkpoint, validate_export_payload,
)
from soma.infrastructure.domain.workbook_normalization import (
    normalize_workbook_row, workbook_logical_fingerprint,
)
from soma.ticket_import.parsing.xlsx_security import (
    XlsxPreflightResult, XlsxResourceLimits, preflight_xlsx,
)

INFRASTRUCTURE_XLSX_LIMITS = XlsxResourceLimits(
    compressed_file_bytes=268_435_456,
    total_expanded_bytes=1_073_741_824,
    single_part_bytes=1_073_741_824,
    zip_entries=20_000,
    expansion_ratio=100,
)


def preflight_infrastructure_workbook(path: str) -> XlsxPreflightResult:
    """Capture and security-check one untrusted source using LLD-08 limits."""
    try:
        return preflight_xlsx(path, limits=INFRASTRUCTURE_XLSX_LIMITS)
    except SomaError as exc:
        if exc.code == "IMPORT_SOURCE_UNAVAILABLE":
            raise SomaError(
                "WORKBOOK_DIRECTORY_UNAVAILABLE", "Infrastructure workbook source is unavailable",
            ) from exc
        if exc.code in ("XLSX_UNSAFE_CONTAINER", "XLSX_RESOURCE_LIMIT", "IMPORT_FILE_UNSTABLE"):
            raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook failed safety preflight") from exc
        raise


@dataclass(frozen=True, slots=True)
class WorkbookProfileSummary:
    mode: str
    source_installation_scope_id: str
    same_installation: bool
    network_element_rows: int
    ip_rows: int
    logical_fingerprint: str


def _profile_error(message: str) -> SomaError:
    return SomaError("WORKBOOK_UNSUPPORTED", message)


def _checked_cell(cell) -> object:
    if cell.data_type == "f":
        raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook contains a formula")
    value = cell.value
    if value is not None and type(value) not in (str, int, bool):
        raise _profile_error("Infrastructure workbook contains an unsupported cell value")
    if isinstance(value, str) and (len(value.encode("utf-8")) > 16_384 or "\x00" in value):
        raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook cell exceeds its safety bound")
    return value


def _checked_header(row, name: str) -> None:
    if row is None:
        raise _profile_error("Infrastructure workbook is missing a required header")
    actual = tuple(_checked_cell(cell) for cell in row)
    expected = HEADERS[name]
    if actual == expected:
        return
    for item in actual[len(expected):]:
        if not isinstance(item, str):
            continue
        lowered = item.casefold().replace("_", "").replace(" ", "")
        if any(token in lowered for token in (
            "password", "privatekey", "secret", "credential", "token",
        )):
            raise SomaError("SECRET_FIELD_FORBIDDEN", "Workbook has a forbidden secret header")
        if any(token in lowered for token in (
            "interface", "port", "topology", "connectivity", "reachability",
        )):
            raise SomaError("TOPOLOGY_FIELD_FORBIDDEN", "Workbook has a forbidden topology header")
    raise _profile_error("Infrastructure workbook header does not match the profile")


def inspect_infrastructure_workbook(
    captured: XlsxPreflightResult, *, current_data_instance_id: str,
) -> WorkbookProfileSummary:
    """Inspect profile structure from the exact immutable preflight snapshot."""
    require_uuid4(current_data_instance_id)
    try:
        workbook = load_workbook(
            captured.semantic_stream(), read_only=True, data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook cannot be parsed") from exc
    try:
        for name in SHEET_ORDER:
            if name not in workbook or workbook[name].sheet_state != "visible":
                raise _profile_error("Infrastructure workbook is missing a visible required sheet")
        counts: dict[str, int] = {}
        metadata: dict[str, object] = {}
        rows = workbook["Metadata"].iter_rows()
        _checked_header(next(rows, None), "Metadata")
        for count, row in enumerate(rows, 1):
            if count > 128:
                raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook metadata row count exceeds its bound")
            values = tuple(_checked_cell(cell) for cell in row)
            if len(values) > 2 and any(value is not None for value in values[2:]):
                raise _profile_error("Infrastructure workbook metadata row exceeds its profile width")
            key = values[0] if values else None
            if not isinstance(key, str) or not key or key in metadata:
                raise _profile_error("Infrastructure workbook metadata key is invalid or duplicate")
            metadata[key] = values[1] if len(values) > 1 else None
        if not set(METADATA_KEYS).issubset(metadata):
            raise _profile_error("Infrastructure workbook is missing required metadata")
        if metadata["FormatId"] != FORMAT_ID or metadata["WorkbookVersion"] != WORKBOOK_VERSION:
            raise _profile_error("Infrastructure workbook format or version is unsupported")
        mode = metadata["Mode"]
        if mode not in MODES:
            raise _profile_error("Infrastructure workbook mode is unsupported")
        source_id = metadata["SourceInstallationScopeId"]
        try:
            require_uuid4(source_id)
        except ValidationError as exc:
            raise _profile_error("Infrastructure workbook source scope is invalid") from exc
        generated_at = metadata["GeneratedAtUtc"]
        if not isinstance(generated_at, str):
            raise _profile_error("Infrastructure workbook generation time is invalid")
        try:
            datetime.strptime(generated_at, "%Y-%m-%dT%H:%M:%SZ")
        except ValueError as exc:
            raise _profile_error("Infrastructure workbook generation time is invalid") from exc
        scope_json = metadata["ExportScopeJson"]
        if not isinstance(scope_json, str):
            raise _profile_error("Infrastructure workbook export scope is invalid")
        try:
            export_scope = validate_value(
                "INFRA_EXPORT_SCOPE_V1", loads_strict(scope_json, max_bytes=16_384),
            )
        except ValidationError as exc:
            raise _profile_error("Infrastructure workbook export scope is invalid") from exc
        same_installation = source_id == current_data_instance_id
        normalized: dict[str, list[str]] = {}
        for name in ("Network Elements", "IP Addresses"):
            sheet = workbook[name]
            rows = sheet.iter_rows()
            _checked_header(next(rows, None), name)
            expected = HEADERS[name]
            maximum = 100_000 if name == "Network Elements" else 500_000
            count = 0
            normalized_rows: list[str] = []
            for row in rows:
                count += 1
                if count > maximum:
                    raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook row count exceeds its bound")
                values = tuple(_checked_cell(cell) for cell in row)
                if len(values) > len(expected) and any(value is not None for value in values[len(expected):]):
                    raise _profile_error("Infrastructure workbook row exceeds its profile width")
                try:
                    _, row_fingerprint = normalize_workbook_row(
                        name, values, same_installation=same_installation,
                    )
                    normalized_rows.append(row_fingerprint)
                except (SomaError, ValidationError) as exc:
                    raise _profile_error("Infrastructure workbook row is invalid") from exc
            counts[name] = count
            normalized[name] = normalized_rows
        logical_fingerprint = workbook_logical_fingerprint(
            mode=mode, source_installation_scope_id=source_id,
            export_scope=export_scope,
            network_rows=normalized["Network Elements"],
            ip_rows=normalized["IP Addresses"],
        )
        return WorkbookProfileSummary(
            mode=mode, source_installation_scope_id=source_id,
            same_installation=same_installation,
            network_element_rows=counts["Network Elements"],
            ip_rows=counts["IP Addresses"],
            logical_fingerprint=logical_fingerprint,
        )
    finally:
        workbook.close()

_GENERATED_PARTS = frozenset({
    "docProps/app.xml", "docProps/core.xml", "xl/theme/theme1.xml",
    "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml",
    "xl/worksheets/sheet3.xml", "xl/styles.xml", "_rels/.rels",
    "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml",
})


def _cell(sheet, value):
    if value is not None and type(value) not in (str, int, bool):
        raise ValidationError("Workbook cells must be text, exact integers or booleans")
    if isinstance(value, str):
        try:
            encoded = value.encode("utf-8", errors="strict")
        except UnicodeError as exc:
            raise ValidationError("Workbook cell text is not valid UTF-8") from exc
        if len(encoded) > 16_384 or "\x00" in value:
            raise ValidationError("Workbook cell text violates the profile bound")
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        # openpyxl otherwise treats leading '=' as an executable formula.
        cell.data_type = "s"
    return cell


def _verify_cell(cell) -> None:
    if cell.data_type == "f":
        raise ValidationError("Workbook contains a formula")
    value = cell.value
    if value is not None and type(value) not in (str, int, bool):
        raise ValidationError("Workbook contains an unsupported cell value")
    if isinstance(value, str):
        try:
            encoded = value.encode("utf-8", errors="strict")
        except UnicodeError as exc:
            raise ValidationError("Workbook cell text is not valid UTF-8") from exc
        if len(encoded) > 16_384 or "\x00" in value:
            raise ValidationError("Workbook cell text violates the profile bound")


def _append_rows(sheet, rows: Iterable[Sequence[object]], *, maximum: int) -> int:
    count = 0
    width = len(HEADERS[sheet.title])
    for row in rows:
        count += 1
        if count > maximum:
            raise ValidationError("Infrastructure workbook row count exceeds its profile")
        if not isinstance(row, (tuple, list)) or len(row) != width:
            raise ValidationError("Infrastructure workbook row does not match its header")
        if sheet.title == "Network Elements":
            for index in (0, 3, 5, 8, 10, 17, 20):
                if row[index] is not None:
                    require_uuid4(row[index])
            for index in (12, 15, 16):
                value = row[index]
                if value is not None and (type(value) is not int or not 1 <= value <= 120):
                    raise ValidationError("Rack values must be exact integers in 1..120")
        else:
            for index in (0, 1):
                if row[index] is not None:
                    require_uuid4(row[index])
            if not isinstance(row[3], str) or not row[3]:
                raise ValidationError("Exported IP address must be text")
            if type(row[4]) is not bool:
                raise ValidationError("Exported IP primary value must be TRUE or FALSE")
        sheet.append([_cell(sheet, value) for value in row])
    return count


def write_profile_workbook(
    stream: BinaryIO,
    *,
    mode: str,
    scope: dict,
    data_instance_id: str,
    generated_at_utc: int,
    network_elements: Iterable[Sequence[object]] = (),
    ip_addresses: Iterable[Sequence[object]] = (),
) -> tuple[int, int]:
    """Write a bounded profile workbook to a caller-owned unpublished stream."""
    if mode not in MODES:
        raise ValidationError("Unsupported Infrastructure workbook mode")
    require_uuid4(data_instance_id)
    if type(generated_at_utc) is not int or generated_at_utc < 0:
        raise ValidationError("Workbook generation timestamp must be a UTC epoch second")
    validated_scope = validate_value("INFRA_EXPORT_SCOPE_V1", scope)
    scope_field = {
        "all": None,
        "customer": "customer_org_id",
        "site": "site_id",
        "network_elements": "network_element_ids",
    }[validated_scope["scope_kind"]]
    present = {key for key, value in validated_scope.items()
               if key != "scope_kind" and value is not None}
    if present != ({scope_field} if scope_field is not None else set()):
        raise ValidationError("Workbook export scope fields disagree with scope_kind")
    if scope_field == "network_element_ids" and not validated_scope[scope_field]:
        raise ValidationError("Network Element export scope cannot be empty")
    try:
        generated_at_text = datetime.fromtimestamp(
            generated_at_utc, timezone.utc
        ).strftime("%Y-%m-%dT%H:%M:%SZ")
    except (OverflowError, OSError, ValueError) as exc:
        raise ValidationError("Workbook generation timestamp is out of range") from exc
    metadata = (
        FORMAT_ID, WORKBOOK_VERSION, mode, data_instance_id,
        generated_at_text,
        canonical_json_bytes(validated_scope).decode("utf-8"),
    )
    workbook = Workbook(write_only=True)
    sheets = {name: workbook.create_sheet(name) for name in SHEET_ORDER}
    for name, sheet in sheets.items():
        sheet.append([_cell(sheet, value) for value in HEADERS[name]])
    metadata_sheet = sheets["Metadata"]
    for key, value in zip(METADATA_KEYS, metadata):
        metadata_sheet.append([_cell(metadata_sheet, key), _cell(metadata_sheet, value)])
    network_count = _append_rows(
        sheets["Network Elements"], network_elements, maximum=100_000,
    )
    ip_count = _append_rows(sheets["IP Addresses"], ip_addresses, maximum=500_000)
    if mode == "registration_template" and (network_count or ip_count):
        raise ValidationError("Registration template must contain only headers")
    workbook.save(stream)
    return network_count, ip_count


def verify_profile_workbook(
    stream: BinaryIO,
    *,
    mode: str,
    data_instance_id: str,
    expected_scope: dict,
    expected_generated_at_utc: int,
    expected_network_rows: int,
    expected_ip_rows: int,
) -> tuple[str, int]:
    """Verify one captured unpublished workbook and return its SHA-256 and size."""
    if mode not in MODES:
        raise ValidationError("Unsupported Infrastructure workbook mode")
    require_uuid4(data_instance_id)
    validated_scope = validate_value("INFRA_EXPORT_SCOPE_V1", expected_scope)
    if type(expected_generated_at_utc) is not int or expected_generated_at_utc < 0:
        raise ValidationError("Expected workbook generation time is invalid")
    try:
        expected_time = datetime.fromtimestamp(
            expected_generated_at_utc, timezone.utc,
        ).strftime("%Y-%m-%dT%H:%M:%SZ")
    except (OverflowError, OSError, ValueError) as exc:
        raise ValidationError("Expected workbook generation time is out of range") from exc
    if any(type(value) is not int or value < 0 for value in
           (expected_network_rows, expected_ip_rows)):
        raise ValidationError("Expected workbook row counts must be nonnegative integers")
    digest = hashlib.sha256()
    size = 0
    with tempfile.SpooledTemporaryFile(max_size=8_388_608) as captured:
        stream.seek(0)
        while chunk := stream.read(1_048_576):
            size += len(chunk)
            if size > 268_435_456:
                raise ValidationError("Workbook exceeds the compressed file bound")
            digest.update(chunk)
            captured.write(chunk)
        if size == 0:
            raise ValidationError("Workbook is empty")
        captured.seek(0)
        try:
            with zipfile.ZipFile(captured) as archive:
                infos = archive.infolist()
                if len(infos) > 20_000 or sum(info.file_size for info in infos) > 1_073_741_824:
                    raise ValidationError("Workbook archive exceeds its resource bounds")
                seen = set()
                for info in infos:
                    name = info.filename
                    parts = name.split("/")
                    key = name.casefold()
                    if (not name or name.startswith("/") or ":" in name or "\\" in name or "\x00" in name
                            or any(part in ("", ".", "..") for part in parts[:-1])
                            or key in seen or posixpath.normpath(name) != name):
                        raise ValidationError("Workbook archive has an unsafe or duplicate part")
                    seen.add(key)
                    lower = name.lower()
                    if any(marker in lower for marker in (
                        "vbaproject", "activex/", "embeddings/", "externallinks/",
                        "oleobject", "customui/",
                    )):
                        raise ValidationError("Workbook contains active or external content")
                    if lower.endswith(".rels") or name == "[Content_Types].xml":
                        if info.file_size > 1_048_576:
                            raise ValidationError("Workbook relationship metadata is too large")
                        raw = archive.read(info)
                        try:
                            root = ElementTree.fromstring(raw)
                        except ElementTree.ParseError as exc:
                            raise ValidationError("Workbook relationship metadata is malformed") from exc
                        for node in root.iter():
                            attributes = {key.rsplit("}", 1)[-1].lower(): value.lower()
                                          for key, value in node.attrib.items()}
                            if attributes.get("targetmode") == "external":
                                raise ValidationError("Workbook contains an external relationship")
                            if any("macroenabled" in value or "vba" in value
                                   for value in attributes.values()):
                                raise ValidationError("Workbook contains macro-enabled content")
                if {info.filename for info in infos} != _GENERATED_PARTS:
                    raise ValidationError("Workbook archive parts differ from generated profile")
            captured.seek(0)
            workbook = load_workbook(captured, read_only=True, data_only=False, keep_links=False)
            try:
                if workbook.sheetnames != list(SHEET_ORDER):
                    raise ValidationError("Workbook sheet order differs from the profile")
                for name in SHEET_ORDER:
                    sheet = workbook[name]
                    if sheet.sheet_state != "visible":
                        raise ValidationError("Workbook has a hidden authoritative sheet")
                    iterator = sheet.iter_rows()
                    header = tuple(cell.value for cell in next(iterator, ()))
                    if header != HEADERS[name]:
                        raise ValidationError("Workbook header differs from the profile")
                    if name == "Metadata":
                        metadata = []
                        for row in iterator:
                            if len(row) != len(HEADERS[name]):
                                raise ValidationError("Workbook metadata row width differs from the profile")
                            for cell in row:
                                _verify_cell(cell)
                            metadata.append(tuple(cell.value for cell in row))
                            if len(metadata) > 128:
                                raise ValidationError("Workbook metadata exceeds its row bound")
                        if tuple(key for key, _ in metadata) != METADATA_KEYS:
                            raise ValidationError("Workbook metadata keys differ from the profile")
                        values = dict(metadata)
                        if (values["FormatId"] != FORMAT_ID
                                or values["WorkbookVersion"] != WORKBOOK_VERSION
                                or values["Mode"] != mode
                                or values["SourceInstallationScopeId"] != data_instance_id):
                            raise ValidationError("Workbook metadata differs from the expected profile")
                        try:
                            export_scope = loads_strict(values["ExportScopeJson"], max_bytes=65_536)
                            parsed_scope = validate_value("INFRA_EXPORT_SCOPE_V1", export_scope)
                        except (TypeError, ValueError) as exc:
                            raise ValidationError("Workbook metadata values are malformed") from exc
                        if (parsed_scope != validated_scope
                                or values["GeneratedAtUtc"] != expected_time):
                            raise ValidationError("Workbook metadata differs from generation inputs")
                        continue
                    count = 0
                    maximum = 100_000 if name == "Network Elements" else 500_000
                    for row in iterator:
                        count += 1
                        if count > maximum:
                            raise ValidationError("Workbook row count exceeds its profile")
                        if len(row) > len(HEADERS[name]):
                            raise ValidationError("Workbook row width exceeds its header")
                        for cell in row:
                            _verify_cell(cell)
                    expected = expected_network_rows if name == "Network Elements" else expected_ip_rows
                    if count != expected:
                        raise ValidationError("Workbook row count differs from verified generation")
            finally:
                workbook.close()
        except (zipfile.BadZipFile, zipfile.LargeZipFile, KeyError,
                ElementTree.ParseError, OSError, EOFError) as exc:
            raise ValidationError("Workbook is not a valid profile XLSX archive") from exc
    return digest.hexdigest(), size



@dataclass(frozen=True, slots=True)
class InfrastructureWorkbookExportResult:
    job_id: str
    export_id: str
    final_filename: str


_EXPORT_JOB_JSON_BYTES = 65_536


def _scope_sql(scope: dict) -> tuple[str, tuple[object, ...]]:
    clauses = ["n.lifecycle_state='active'"]
    values: list[object] = []
    kind = scope["scope_kind"]
    if kind == "customer":
        clauses.append("s.customer_org_id=?")
        values.append(scope["customer_org_id"])
    elif kind == "site":
        clauses.append("n.site_id=?")
        values.append(scope["site_id"])
    elif kind == "network_elements":
        identities = tuple(scope["network_element_ids"])
        if not identities:
            raise IntegrityFailure("persisted workbook export scope is empty")
        clauses.append(
            "n.network_element_id IN (" + ",".join("?" for _ in identities) + ")"
        )
        values.extend(identities)
    elif kind != "all":
        raise IntegrityFailure("persisted workbook export scope kind is invalid")
    return " AND ".join(clauses), tuple(values)


def _network_export_rows(snapshot: ReadSnapshot, scope: dict):
    where, values = _scope_sql(scope)
    cursor = snapshot.connection.execute(
        """
        SELECT
            n.network_element_id,n.operational_name,n.manufacturer_serial,
            ma.network_element_model_id,m.name,
            s.site_id,s.name,s.address_text,
            rm.room_id,rm.name,
            r.rack_id,r.name,r.height_u,r.row_label,r.column_label,
            p.u_start,p.u_span,
            ca.cloud_deployment_id,ct.name,cd.name,
            cc.parent_network_element_id
        FROM network_elements n
        JOIN sites s ON s.site_id=n.site_id
        JOIN network_element_placement_current p
          ON p.network_element_id=n.network_element_id
        LEFT JOIN racks r ON r.rack_id=p.rack_id
        LEFT JOIN rooms rm ON rm.room_id=r.room_id
        LEFT JOIN network_element_model_current ma
          ON ma.network_element_id=n.network_element_id
        LEFT JOIN network_element_models m
          ON m.network_element_model_id=ma.network_element_model_id
        LEFT JOIN cloud_assignment_current ca
          ON ca.network_element_id=n.network_element_id
        LEFT JOIN cloud_deployments cd
          ON cd.cloud_deployment_id=ca.cloud_deployment_id
        LEFT JOIN cloud_types ct ON ct.cloud_type_id=cd.cloud_type_id
        LEFT JOIN network_element_containment_current cc
          ON cc.child_network_element_id=n.network_element_id
        WHERE """ + where + """
        ORDER BY n.network_element_id
        """,
        values,
    )
    for row in cursor:
        yield (*tuple(row), None)


def _ip_export_rows(snapshot: ReadSnapshot, scope: dict):
    where, values = _scope_sql(scope)
    cursor = snapshot.connection.execute(
        """
        SELECT ip.network_element_ip_id,ip.network_element_id,n.operational_name,
               ip.canonical_address,ip.is_primary
        FROM network_element_ip_current ip
        JOIN network_elements n ON n.network_element_id=ip.network_element_id
        JOIN sites s ON s.site_id=n.site_id
        WHERE ip.active=1 AND """ + where + """
        ORDER BY ip.network_element_id,ip.network_element_ip_id
        """,
        values,
    )
    for row in cursor:
        yield (row[0], row[1], row[2], row[3], bool(row[4]))


def _scoped_export_counts(snapshot: ReadSnapshot, scope: dict) -> tuple[int, int]:
    where, values = _scope_sql(scope)
    network_count = int(snapshot.connection.execute(
        """
        SELECT count(*)
        FROM network_elements n
        JOIN sites s ON s.site_id=n.site_id
        WHERE """ + where,
        values,
    ).fetchone()[0])
    ip_count = int(snapshot.connection.execute(
        """
        SELECT count(*)
        FROM network_element_ip_current ip
        JOIN network_elements n ON n.network_element_id=ip.network_element_id
        JOIN sites s ON s.site_id=n.site_id
        WHERE ip.active=1 AND """ + where,
        values,
    ).fetchone()[0])
    return network_count, ip_count


class InfrastructureWorkbookExportWorker:
    """Crash-recoverable Infrastructure workbook export and evidence publisher."""

    def __init__(
        self,
        connection_factory: ConnectionFactory,
        *,
        clock=None,
    ) -> None:
        self._factory = connection_factory
        coordinator_kwargs = {} if clock is None else {"clock": clock}
        self._jobs = DurableJobCoordinator(
            connection_factory,
            JobTypeRegistry(INFRASTRUCTURE_JOB_CONTRACTS),
            **coordinator_kwargs,
        )

    @staticmethod
    def _payload(claim: DurableJobClaim) -> dict:
        if claim.job_type != EXPORT_JOB_TYPE or claim.contract_version != 1:
            raise ValidationError("claim is not INFRA_WORKBOOK_EXPORT_V1")
        value = loads_canonical_json(
            claim.payload_json,
            max_bytes=_EXPORT_JOB_JSON_BYTES,
            max_depth=8,
            max_collection_items=512,
        )
        if not isinstance(value, dict):
            raise IntegrityFailure("persisted Infrastructure export payload is not an object")
        try:
            validate_export_payload(value)
        except ValidationError as exc:
            raise IntegrityFailure(
                "persisted Infrastructure export payload violates its contract"
            ) from exc
        return value

    @staticmethod
    def _checkpoint(claim: DurableJobClaim) -> dict | None:
        if claim.checkpoint_json is None:
            return None
        value = loads_canonical_json(
            claim.checkpoint_json,
            max_bytes=_EXPORT_JOB_JSON_BYTES,
            max_depth=4,
            max_collection_items=32,
        )
        if not isinstance(value, dict):
            raise IntegrityFailure("persisted Infrastructure export checkpoint is not an object")
        try:
            validate_export_checkpoint(value)
        except ValidationError as exc:
            raise IntegrityFailure(
                "persisted Infrastructure export checkpoint violates its contract"
            ) from exc
        return value

    def _created_at(self, claim: DurableJobClaim) -> int:
        with ReadSnapshot(self._factory) as snapshot:
            row = snapshot.connection.execute(
                "SELECT created_at_utc FROM durable_jobs WHERE job_id=?",
                (claim.job_id,),
            ).fetchone()
        if row is None or type(row[0]) is not int or row[0] < 0:
            raise IntegrityFailure("Infrastructure export job creation time is unavailable")
        return int(row[0])

    @staticmethod
    def _checkpoint_value(
        *,
        phase: str,
        export_id: str,
        temp_filename: str | None,
        verified_sha256: str | None = None,
        verified_size_bytes: int | None = None,
        final_filename: str | None = None,
    ) -> dict:
        return {
            "phase": phase,
            "export_id": export_id,
            "temp_filename": temp_filename,
            "verified_sha256": verified_sha256,
            "verified_size_bytes": verified_size_bytes,
            "final_filename": final_filename,
        }

    @staticmethod
    def _temp_filename(claim: DurableJobClaim, export_id: str) -> str:
        return f".soma-{export_id}-{claim.job_id}.tmp"

    @staticmethod
    def _require_filename_identity(
        claim: DurableJobClaim,
        checkpoint: dict,
        *,
        generated_at_utc: int,
        payload: dict,
    ) -> tuple[str, str]:
        export_id = str(checkpoint["export_id"])
        expected_temp = InfrastructureWorkbookExportWorker._temp_filename(claim, export_id)
        expected_final = artifact_filename(payload["mode"], export_id, generated_at_utc)
        if checkpoint["temp_filename"] not in (None, expected_temp):
            raise IntegrityFailure("Infrastructure export temp identity changed")
        if checkpoint["final_filename"] not in (None, expected_final):
            raise IntegrityFailure("Infrastructure export final identity changed")
        return expected_temp, expected_final

    @staticmethod
    def _destination(payload: dict) -> Path:
        destination = Path(str(payload["destination_directory"]))
        try:
            if not destination.is_dir():
                raise SomaError(
                    "WORKBOOK_DIRECTORY_UNAVAILABLE",
                    "Infrastructure workbook destination is unavailable",
                )
        except OSError as exc:
            raise SomaError(
                "WORKBOOK_DIRECTORY_UNAVAILABLE",
                "Infrastructure workbook destination is unavailable",
            ) from exc
        return destination

    @staticmethod
    def _unlink_owned(path: Path) -> None:
        try:
            path.unlink(missing_ok=True)
        except OSError as exc:
            raise SomaError(
                "WORKBOOK_DIRECTORY_UNAVAILABLE",
                "Infrastructure workbook temporary artifact cannot be removed",
            ) from exc

    @staticmethod
    def _same_file(left: Path, right: Path) -> bool:
        try:
            return os.path.samefile(left, right)
        except OSError:
            return False

    def _claim_was_cancelled(self, claim: DurableJobClaim) -> bool:
        with ReadSnapshot(self._factory) as snapshot:
            row = snapshot.connection.execute(
                "SELECT state,job_type,contract_version,attempt_count "
                "FROM durable_jobs WHERE job_id=?",
                (claim.job_id,),
            ).fetchone()
            if (
                row is None
                or str(row[0]) != "cancelled"
                or str(row[1]) != claim.job_type
                or int(row[2]) != claim.contract_version
                or int(row[3]) != claim.attempt_ordinal
            ):
                return False
            attempt = snapshot.connection.execute(
                "SELECT run_id,outcome FROM job_attempts "
                "WHERE job_id=? AND ordinal=?",
                (claim.job_id, claim.attempt_ordinal),
            ).fetchone()
        return attempt is not None and tuple(attempt) == (claim.run_id, "cancelled")

    def _require_data_instance(self, reader, payload: dict) -> None:
        if DataInstanceIdentityReader.get(reader) != payload["data_instance_id"]:
            raise SomaError(
                "WORKBOOK_STALE",
                "Infrastructure export belongs to a different data instance",
            )

    def _write_temp(
        self,
        path: Path,
        *,
        payload: dict,
        generated_at_utc: int,
    ) -> None:
        self._unlink_owned(path)
        try:
            with ReadSnapshot(self._factory) as snapshot:
                self._require_data_instance(snapshot, payload)
                if payload["mode"] == "registration_template":
                    expected_network, expected_ip = 0, 0
                    network_rows = ()
                    ip_rows = ()
                else:
                    expected_network, expected_ip = _scoped_export_counts(
                        snapshot, payload["scope"]
                    )
                    network_rows = _network_export_rows(snapshot, payload["scope"])
                    ip_rows = _ip_export_rows(snapshot, payload["scope"])
                with path.open("xb") as stream:
                    actual_network, actual_ip = write_profile_workbook(
                        stream,
                        mode=payload["mode"],
                        scope=payload["scope"],
                        data_instance_id=payload["data_instance_id"],
                        generated_at_utc=generated_at_utc,
                        network_elements=network_rows,
                        ip_addresses=ip_rows,
                    )
                if (actual_network, actual_ip) != (expected_network, expected_ip):
                    raise IntegrityFailure(
                        "Infrastructure export snapshot row counts changed during generation"
                    )
        except Exception:
            self._unlink_owned(path)
            raise

    def _verify_artifact(
        self,
        path: Path,
        *,
        payload: dict,
        generated_at_utc: int,
        expected_sha256: str | None = None,
        expected_size_bytes: int | None = None,
    ) -> tuple[str, int]:
        try:
            captured = preflight_infrastructure_workbook(str(path))
        except OSError as exc:
            raise SomaError(
                "WORKBOOK_DIRECTORY_UNAVAILABLE",
                "Infrastructure workbook artifact cannot be reopened",
            ) from exc
        try:
            summary = inspect_infrastructure_workbook(
                captured,
                current_data_instance_id=payload["data_instance_id"],
            )
        finally:
            captured.close()
        if (
            summary.mode != payload["mode"]
            or summary.source_installation_scope_id != payload["data_instance_id"]
            or not summary.same_installation
        ):
            raise IntegrityFailure("Infrastructure export artifact identity changed")
        try:
            with path.open("rb") as stream:
                digest, size = verify_profile_workbook(
                    stream,
                    mode=payload["mode"],
                    data_instance_id=payload["data_instance_id"],
                    expected_scope=payload["scope"],
                    expected_generated_at_utc=generated_at_utc,
                    expected_network_rows=summary.network_element_rows,
                    expected_ip_rows=summary.ip_rows,
                )
        except OSError as exc:
            raise SomaError(
                "WORKBOOK_DIRECTORY_UNAVAILABLE",
                "Infrastructure workbook artifact cannot be reopened",
            ) from exc
        if expected_sha256 is not None and digest != expected_sha256:
            raise IntegrityFailure("published Infrastructure workbook hash changed")
        if expected_size_bytes is not None and size != expected_size_bytes:
            raise IntegrityFailure("published Infrastructure workbook size changed")
        return digest, size

    def _evidence(
        self,
        uow: UnitOfWork,
        claim: DurableJobClaim,
        *,
        payload: dict,
        checkpoint: dict,
        generated_at_utc: int,
    ) -> None:
        self._require_data_instance(uow, payload)
        export_id = str(checkpoint["export_id"])
        expected = (
            payload["mode"],
            payload["data_instance_id"],
            canonical_json_bytes(payload["scope"]).decode("utf-8"),
            generated_at_utc,
            checkpoint["final_filename"],
            checkpoint["verified_sha256"],
            checkpoint["verified_size_bytes"],
            payload["export_request_id"],
        )
        row = uow.connection.execute(
            """
            SELECT mode,installation_scope_id,filter_scope_json,generated_at_utc,
                   artifact_filename,artifact_sha256,artifact_size_bytes,command_id
            FROM infrastructure_workbook_exports
            WHERE export_id=?
            """,
            (export_id,),
        ).fetchone()
        if row is None:
            uow.connection.execute(
                """
                INSERT INTO infrastructure_workbook_exports(
                    export_id,mode,installation_scope_id,filter_scope_json,
                    generated_at_utc,artifact_filename,artifact_sha256,
                    artifact_size_bytes,command_id
                ) VALUES (?,?,?,?,?,?,?,?,?)
                """,
                (export_id, *expected),
            )
        elif tuple(row) != expected:
            raise IntegrityFailure("Infrastructure workbook export evidence changed")

    def _finish_published(
        self,
        claim: DurableJobClaim,
        *,
        payload: dict,
        checkpoint: dict,
        generated_at_utc: int,
    ) -> None:
        published = {
            **checkpoint,
            "phase": "published",
        }
        recorded = {
            **published,
            "phase": "evidence_recorded",
        }
        with UnitOfWork(self._factory) as uow:
            self._jobs.assert_claim_current(uow, claim)
            self._jobs.checkpoint_in_uow(uow, claim, published)
            self._evidence(
                uow,
                claim,
                payload=payload,
                checkpoint=published,
                generated_at_utc=generated_at_utc,
            )
            self._jobs.checkpoint_in_uow(uow, claim, recorded)
            self._jobs.complete_in_uow(uow, claim)

    def _publish_from_verified(
        self,
        claim: DurableJobClaim,
        *,
        payload: dict,
        checkpoint: dict,
        generated_at_utc: int,
        destination: Path,
    ) -> None:
        temp = destination / str(checkpoint["temp_filename"])
        final = destination / str(checkpoint["final_filename"])
        expected_sha = str(checkpoint["verified_sha256"])
        expected_size = int(checkpoint["verified_size_bytes"])

        if final.exists():
            if not temp.exists() or not self._same_file(temp, final):
                raise IntegrityFailure(
                    "Infrastructure export final artifact exists without job ownership proof"
                )
            self._verify_artifact(
                final,
                payload=payload,
                generated_at_utc=generated_at_utc,
                expected_sha256=expected_sha,
                expected_size_bytes=expected_size,
            )
            self._finish_published(
                claim,
                payload=payload,
                checkpoint=checkpoint,
                generated_at_utc=generated_at_utc,
            )
            self._unlink_owned(temp)
            return

        self._verify_artifact(
            temp,
            payload=payload,
            generated_at_utc=generated_at_utc,
            expected_sha256=expected_sha,
            expected_size_bytes=expected_size,
        )
        recorded = {**checkpoint, "phase": "evidence_recorded"}
        published = {**checkpoint, "phase": "published"}
        with UnitOfWork(self._factory) as uow:
            self._jobs.assert_claim_current(uow, claim)
            self._require_data_instance(uow, payload)
            try:
                os.link(temp, final)
            except FileExistsError as exc:
                raise IntegrityFailure(
                    "Infrastructure export destination filename collision"
                ) from exc
            except OSError as exc:
                raise SomaError(
                    "WORKBOOK_DIRECTORY_UNAVAILABLE",
                    "Infrastructure workbook artifact cannot be atomically published",
                ) from exc
            self._jobs.checkpoint_in_uow(uow, claim, published)
            self._evidence(
                uow,
                claim,
                payload=payload,
                checkpoint=published,
                generated_at_utc=generated_at_utc,
            )
            self._jobs.checkpoint_in_uow(uow, claim, recorded)
            self._jobs.complete_in_uow(uow, claim)
        self._unlink_owned(temp)

    def run(self, claim: DurableJobClaim) -> InfrastructureWorkbookExportResult:
        payload = self._payload(claim)
        checkpoint = self._checkpoint(claim)
        generated_at_utc = self._created_at(claim)
        destination = self._destination(payload)

        if checkpoint is None:
            checkpoint = self._checkpoint_value(
                phase="queued",
                export_id=new_uuid4(),
                temp_filename=None,
            )
            self._jobs.checkpoint(claim, checkpoint)

        if checkpoint["export_id"] is None:
            raise IntegrityFailure("Infrastructure export checkpoint lacks its export identity")
        expected_temp, expected_final = self._require_filename_identity(
            claim,
            checkpoint,
            generated_at_utc=generated_at_utc,
            payload=payload,
        )
        phase = str(checkpoint["phase"])

        if phase == "queued":
            checkpoint = self._checkpoint_value(
                phase="writing",
                export_id=str(checkpoint["export_id"]),
                temp_filename=expected_temp,
            )
            self._jobs.checkpoint(claim, checkpoint)
            phase = "writing"

        if phase == "writing":
            temp = destination / expected_temp
            self._write_temp(
                temp,
                payload=payload,
                generated_at_utc=generated_at_utc,
            )
            checkpoint = self._checkpoint_value(
                phase="verifying",
                export_id=str(checkpoint["export_id"]),
                temp_filename=expected_temp,
            )
            try:
                self._jobs.checkpoint(claim, checkpoint)
            except JobClaimConflict:
                self._unlink_owned(temp)
                raise
            phase = "verifying"

        if phase == "verifying":
            temp = destination / expected_temp
            if not temp.is_file():
                checkpoint = self._checkpoint_value(
                    phase="writing",
                    export_id=str(checkpoint["export_id"]),
                    temp_filename=expected_temp,
                )
                self._jobs.checkpoint(claim, checkpoint)
                self._write_temp(
                    temp,
                    payload=payload,
                    generated_at_utc=generated_at_utc,
                )
            digest, size = self._verify_artifact(
                temp,
                payload=payload,
                generated_at_utc=generated_at_utc,
            )
            checkpoint = self._checkpoint_value(
                phase="verified",
                export_id=str(checkpoint["export_id"]),
                temp_filename=expected_temp,
                verified_sha256=digest,
                verified_size_bytes=size,
                final_filename=expected_final,
            )
            try:
                self._jobs.checkpoint(claim, checkpoint)
            except JobClaimConflict:
                self._unlink_owned(temp)
                raise
            phase = "verified"

        if phase == "verified":
            try:
                self._publish_from_verified(
                    claim,
                    payload=payload,
                    checkpoint=checkpoint,
                    generated_at_utc=generated_at_utc,
                    destination=destination,
                )
            except JobClaimConflict:
                # A verified temp belongs to a recoverable newer attempt unless
                # the exact attempt was durably cancelled. Cancellation alone
                # authorizes deletion of this unpublished job-owned artifact.
                if self._claim_was_cancelled(claim):
                    self._unlink_owned(destination / expected_temp)
                raise
            return InfrastructureWorkbookExportResult(
                claim.job_id,
                str(checkpoint["export_id"]),
                expected_final,
            )

        if phase == "published":
            final = destination / expected_final
            self._verify_artifact(
                final,
                payload=payload,
                generated_at_utc=generated_at_utc,
                expected_sha256=str(checkpoint["verified_sha256"]),
                expected_size_bytes=int(checkpoint["verified_size_bytes"]),
            )
            self._finish_published(
                claim,
                payload=payload,
                checkpoint=checkpoint,
                generated_at_utc=generated_at_utc,
            )
            return InfrastructureWorkbookExportResult(
                claim.job_id,
                str(checkpoint["export_id"]),
                expected_final,
            )

        if phase == "evidence_recorded":
            with UnitOfWork(self._factory) as uow:
                self._jobs.assert_claim_current(uow, claim)
                self._evidence(
                    uow,
                    claim,
                    payload=payload,
                    checkpoint=checkpoint,
                    generated_at_utc=generated_at_utc,
                )
                self._jobs.complete_in_uow(uow, claim)
            return InfrastructureWorkbookExportResult(
                claim.job_id,
                str(checkpoint["export_id"]),
                expected_final,
            )

        raise IntegrityFailure("Infrastructure export checkpoint phase is unsupported")


def run_export(
    claim: DurableJobClaim,
    connection_factory: ConnectionFactory,
) -> InfrastructureWorkbookExportResult:
    return InfrastructureWorkbookExportWorker(connection_factory).run(claim)
