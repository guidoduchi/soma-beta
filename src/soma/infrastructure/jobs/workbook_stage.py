from __future__ import annotations

from dataclasses import dataclass
import hashlib
import os
from pathlib import Path

from openpyxl import load_workbook

from soma.foundation.contracts.foundation import DurableJobClaim
from soma.foundation.errors import IntegrityFailure, SomaError, ValidationError
from soma.foundation.identifiers import new_uuid4, utc_epoch_seconds
from soma.foundation.jobs import DurableJobCoordinator, JobTypeRegistry
from soma.foundation.persistence.connections import ConnectionFactory
from soma.foundation.persistence.uow import ReadSnapshot, UnitOfWork
from soma.foundation.queries.data_instance_identity import DataInstanceIdentityReader
from soma.foundation.strict_json import canonical_json_bytes, loads_canonical_json
from soma.infrastructure.domain.workbook_normalization import normalize_workbook_row
from soma.infrastructure.domain.workbooks import HEADERS
from soma.infrastructure.jobs import (
    INFRASTRUCTURE_JOB_CONTRACTS,
    STAGE_JOB_TYPE,
    validate_stage_checkpoint,
    validate_stage_payload,
)
from soma.infrastructure.jobs.infrastructure_workbooks import (
    _checked_cell,
    _checked_header,
    inspect_infrastructure_workbook,
    preflight_infrastructure_workbook,
)
from soma.infrastructure.jobs.workbook_proposals import build_workbook_proposal
from soma.infrastructure.settings import (
    IMPORT_DIRECTORY_KEY,
    observe_import_directory,
    require_saved_import_directory,
    validate_import_directory,
)


_STAGE_JOB_JSON_BYTES = 65_536
_MAX_CANDIDATES = 512
_STAGE_BATCH = 250
_PROPOSAL_BATCH = 250


@dataclass(frozen=True, slots=True)
class WorkbookCandidate:
    filename: str
    size_bytes: int
    mtime_ns: int


@dataclass(frozen=True, slots=True)
class InfrastructureWorkbookStageResult:
    job_id: str
    published_run_ids: tuple[str, ...]


def _filename(value: str) -> str:
    if not isinstance(value, str) or not value or value in (".", ".."):
        raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook candidate filename is invalid")
    try:
        encoded = value.encode("utf-8", errors="strict")
    except UnicodeError as exc:
        raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook candidate filename is invalid") from exc
    if len(encoded) > 1024 or any(ch in value for ch in ("/", "\\", ":", "\x00")):
        raise SomaError("WORKBOOK_UNSAFE", "Infrastructure workbook candidate filename is invalid")
    return value


def discover_workbook_candidates(directory: Path) -> tuple[WorkbookCandidate, ...]:
    candidates: list[WorkbookCandidate] = []
    try:
        with os.scandir(directory) as entries:
            for entry in entries:
                if entry.is_symlink() or not entry.is_file(follow_symlinks=False):
                    continue
                if not entry.name.lower().endswith(".xlsx"):
                    continue
                name = _filename(entry.name)
                stat = entry.stat(follow_symlinks=False)
                if stat.st_size < 0 or stat.st_mtime_ns < 0:
                    raise SomaError(
                        "WORKBOOK_UNSAFE",
                        "Infrastructure workbook candidate metadata is invalid",
                    )
                candidates.append(
                    WorkbookCandidate(name, int(stat.st_size), int(stat.st_mtime_ns))
                )
                if len(candidates) > _MAX_CANDIDATES:
                    raise SomaError(
                        "WORKBOOK_UNSAFE",
                        "Infrastructure workbook candidate count exceeds its bounded job profile",
                    )
    except SomaError:
        raise
    except OSError as exc:
        raise SomaError(
            "WORKBOOK_DIRECTORY_UNAVAILABLE",
            "Infrastructure workbook import directory cannot be enumerated",
        ) from exc

    candidates.sort(key=lambda item: (item.filename.casefold(), item.filename))
    folded = [item.filename.casefold() for item in candidates]
    if len(folded) != len(set(folded)):
        raise SomaError(
            "WORKBOOK_UNSAFE",
            "Infrastructure workbook candidate filenames are ambiguous by case",
        )
    return tuple(candidates)


def candidate_manifest_sha256(candidates: tuple[WorkbookCandidate, ...]) -> str:
    return hashlib.sha256(
        canonical_json_bytes(
            {
                "version": "INFRA_WORKBOOK_CANDIDATE_MANIFEST_V1",
                "candidates": [
                    {
                        "filename": item.filename,
                        "size_bytes": item.size_bytes,
                        "mtime_ns": item.mtime_ns,
                    }
                    for item in candidates
                ],
            }
        )
    ).hexdigest()


def _empty_checkpoint(
    *,
    manifest_sha256: str,
    candidate_count: int,
    candidate_index: int = 0,
    published_run_ids: list[str] | None = None,
    phase: str = "discovering",
) -> dict:
    return {
        "phase": phase,
        "candidate_manifest_sha256": manifest_sha256,
        "candidate_count": candidate_count,
        "candidate_index": candidate_index,
        "current_filename": None,
        "current_file_size_bytes": None,
        "current_file_mtime_ns": None,
        "current_file_sha256": None,
        "current_workbook_run_id": None,
        "last_committed_network_element_row": None,
        "last_committed_ip_row": None,
        "published_run_ids": list(published_run_ids or ()),
    }


class InfrastructureWorkbookStageWorker:
    """Durable non-authoritative workbook discovery, staging and review publication."""

    def __init__(
        self,
        connection_factory: ConnectionFactory,
        *,
        setting_service,
        clock=utc_epoch_seconds,
    ) -> None:
        self._factory = connection_factory
        self._setting_service = setting_service
        self._clock = clock
        self._jobs = DurableJobCoordinator(
            connection_factory,
            JobTypeRegistry(INFRASTRUCTURE_JOB_CONTRACTS),
            clock=clock,
        )

    @staticmethod
    def _payload(claim: DurableJobClaim) -> dict:
        if claim.job_type != STAGE_JOB_TYPE or claim.contract_version != 1:
            raise ValidationError("claim is not INFRA_WORKBOOK_STAGE_V1")
        value = loads_canonical_json(
            claim.payload_json,
            max_bytes=_STAGE_JOB_JSON_BYTES,
            max_depth=4,
            max_collection_items=16,
        )
        if not isinstance(value, dict):
            raise IntegrityFailure("persisted Infrastructure stage payload is not an object")
        try:
            validate_stage_payload(value)
        except ValidationError as exc:
            raise IntegrityFailure(
                "persisted Infrastructure stage payload violates its contract"
            ) from exc
        return value

    @staticmethod
    def _checkpoint(claim: DurableJobClaim) -> dict | None:
        if claim.checkpoint_json is None:
            return None
        value = loads_canonical_json(
            claim.checkpoint_json,
            max_bytes=_STAGE_JOB_JSON_BYTES,
            max_depth=5,
            max_collection_items=520,
        )
        if not isinstance(value, dict):
            raise IntegrityFailure("persisted Infrastructure stage checkpoint is not an object")
        try:
            validate_stage_checkpoint(value)
        except ValidationError as exc:
            raise IntegrityFailure(
                "persisted Infrastructure stage checkpoint violates its contract"
            ) from exc
        return value

    def _context(self, reader, payload: dict) -> Path:
        if DataInstanceIdentityReader.get(reader) != payload["data_instance_id"]:
            raise SomaError(
                "WORKBOOK_STALE",
                "Infrastructure workbook check belongs to a different data instance",
            )
        setting = self._setting_service.get_in_reader(reader, IMPORT_DIRECTORY_KEY)
        saved = require_saved_import_directory(setting, payload["setting_revision"])
        expected = validate_import_directory({"path": payload["import_directory"]})
        if saved != expected:
            raise SomaError(
                "WORKBOOK_STALE",
                "Infrastructure import directory changed after job enqueue",
            )
        return observe_import_directory(saved)

    def _discover(self, payload: dict) -> tuple[Path, tuple[WorkbookCandidate, ...], str]:
        with ReadSnapshot(self._factory) as snapshot:
            directory = self._context(snapshot, payload)
        candidates = discover_workbook_candidates(directory)
        return directory, candidates, candidate_manifest_sha256(candidates)

    @staticmethod
    def _candidate_path(directory: Path, candidate: WorkbookCandidate) -> Path:
        path = directory / candidate.filename
        if path.parent != directory:
            raise IntegrityFailure("Infrastructure workbook candidate escaped its directory")
        return path

    @staticmethod
    def _stat_exact(path: Path, expected: WorkbookCandidate) -> os.stat_result:
        try:
            stat = path.stat(follow_symlinks=False)
        except OSError as exc:
            raise SomaError(
                "WORKBOOK_DIRECTORY_UNAVAILABLE",
                "Infrastructure workbook candidate is unavailable",
            ) from exc
        if (
            not path.is_file()
            or int(stat.st_size) != expected.size_bytes
            or int(stat.st_mtime_ns) != expected.mtime_ns
        ):
            raise SomaError(
                "WORKBOOK_STALE",
                "Infrastructure workbook candidate changed during staging",
            )
        return stat

    def _capture(self, directory: Path, candidate: WorkbookCandidate):
        path = self._candidate_path(directory, candidate)
        self._stat_exact(path, candidate)
        captured = preflight_infrastructure_workbook(str(path))
        try:
            self._stat_exact(path, candidate)
            if captured.compressed_file_bytes != candidate.size_bytes:
                raise SomaError(
                    "WORKBOOK_STALE",
                    "Infrastructure workbook candidate size changed during capture",
                )
            return captured
        except Exception:
            captured.close()
            raise

    @staticmethod
    def _iter_normalized(captured, *, same_installation: bool):
        try:
            workbook = load_workbook(
                captured.semantic_stream(),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except Exception as exc:
            raise SomaError(
                "WORKBOOK_UNSAFE",
                "Infrastructure workbook cannot be reopened for bounded staging",
            ) from exc
        try:
            for sheet_name, sheet_kind in (
                ("Network Elements", "network_elements"),
                ("IP Addresses", "ip_addresses"),
            ):
                sheet = workbook[sheet_name]
                rows = sheet.iter_rows()
                _checked_header(next(rows, None), sheet_name)
                for ordinal, row in enumerate(rows, 2):
                    values = tuple(_checked_cell(cell) for cell in row)
                    normalized, fingerprint = normalize_workbook_row(
                        sheet_name,
                        values,
                        same_installation=same_installation,
                    )
                    yield sheet_kind, ordinal, normalized, fingerprint
        finally:
            workbook.close()

    def _start_run(
        self,
        claim: DurableJobClaim,
        checkpoint: dict,
        *,
        candidate: WorkbookCandidate,
        file_sha256: str,
        summary,
    ) -> tuple[str, dict]:
        run_id = new_uuid4()
        relation = "same_installation" if summary.same_installation else "foreign_installation"
        current = {
            **checkpoint,
            "phase": "parsing",
            "current_filename": candidate.filename,
            "current_file_size_bytes": candidate.size_bytes,
            "current_file_mtime_ns": candidate.mtime_ns,
            "current_file_sha256": file_sha256,
            "current_workbook_run_id": run_id,
            "last_committed_network_element_row": None,
            "last_committed_ip_row": None,
        }
        now = int(self._clock())
        with UnitOfWork(self._factory) as uow:
            self._jobs.assert_claim_current(uow, claim)
            uow.connection.execute(
                """
                INSERT INTO infrastructure_workbook_runs(
                    workbook_run_id,source_filename,file_sha256,logical_fingerprint,
                    workbook_version,workbook_mode,source_installation_scope_id,
                    installation_relation,state,network_element_row_count,ip_row_count,
                    warning_count,captured_at_utc,revision
                ) VALUES (?,?,?,?,?,?,?,?, 'validating',?,?,?,?,1)
                """,
                (
                    run_id,
                    candidate.filename,
                    file_sha256,
                    summary.logical_fingerprint,
                    "1.0",
                    summary.mode,
                    summary.source_installation_scope_id,
                    relation,
                    summary.network_element_rows,
                    summary.ip_rows,
                    0,
                    now,
                ),
            )
            self._jobs.checkpoint_in_uow(uow, claim, current)
        return run_id, current

    def _existing_review_run(
        self,
        candidate: WorkbookCandidate,
        *,
        file_sha256: str,
        logical_fingerprint: str,
    ) -> str | None:
        with ReadSnapshot(self._factory) as snapshot:
            rows = snapshot.connection.execute(
                """
                SELECT workbook_run_id FROM infrastructure_workbook_runs
                WHERE source_filename=? AND file_sha256=? AND logical_fingerprint=?
                  AND state IN ('staged','reviewed','accepted','rejected')
                ORDER BY workbook_run_id LIMIT 2
                """,
                (candidate.filename, file_sha256, logical_fingerprint),
            ).fetchall()
        if len(rows) > 1:
            raise IntegrityFailure(
                "multiple published Infrastructure workbook runs share one exact source identity"
            )
        return None if not rows else str(rows[0][0])

    def _require_run_identity(
        self,
        checkpoint: dict,
        candidate: WorkbookCandidate,
        *,
        file_sha256: str,
        summary,
    ) -> str:
        run_id = checkpoint["current_workbook_run_id"]
        if run_id is None:
            raise IntegrityFailure("Infrastructure stage checkpoint lacks current run identity")
        if (
            checkpoint["current_filename"] != candidate.filename
            or checkpoint["current_file_size_bytes"] != candidate.size_bytes
            or checkpoint["current_file_mtime_ns"] != candidate.mtime_ns
            or checkpoint["current_file_sha256"] != file_sha256
        ):
            raise SomaError(
                "WORKBOOK_STALE",
                "Infrastructure workbook source identity changed after checkpoint",
            )
        with ReadSnapshot(self._factory) as snapshot:
            row = snapshot.connection.execute(
                """
                SELECT source_filename,file_sha256,logical_fingerprint,workbook_mode,
                       source_installation_scope_id,installation_relation,state,
                       network_element_row_count,ip_row_count
                FROM infrastructure_workbook_runs WHERE workbook_run_id=?
                """,
                (run_id,),
            ).fetchone()
        expected = (
            candidate.filename,
            file_sha256,
            summary.logical_fingerprint,
            summary.mode,
            summary.source_installation_scope_id,
            "same_installation" if summary.same_installation else "foreign_installation",
            "validating",
            summary.network_element_rows,
            summary.ip_rows,
        )
        if row is None or tuple(row) != expected:
            raise IntegrityFailure(
                "Infrastructure validating run disagrees with its durable source checkpoint"
            )
        return str(run_id)

    def _stage_rows(
        self,
        claim: DurableJobClaim,
        checkpoint: dict,
        captured,
        *,
        same_installation: bool,
    ) -> dict:
        run_id = str(checkpoint["current_workbook_run_id"])
        cursors = {
            "network_elements": checkpoint["last_committed_network_element_row"],
            "ip_addresses": checkpoint["last_committed_ip_row"],
        }
        batch: list[tuple[str, int, dict, str]] = []

        def commit(rows_to_commit: list[tuple[str, int, dict, str]]) -> None:
            nonlocal checkpoint
            if not rows_to_commit:
                return
            with UnitOfWork(self._factory) as uow:
                self._jobs.assert_claim_current(uow, claim)
                state = uow.connection.execute(
                    "SELECT state FROM infrastructure_workbook_runs WHERE workbook_run_id=?",
                    (run_id,),
                ).fetchone()
                if state != ("validating",):
                    raise IntegrityFailure("Infrastructure staging run is no longer validating")
                for sheet_kind, ordinal, normalized, fingerprint in rows_to_commit:
                    encoded = canonical_json_bytes(normalized).decode("utf-8")
                    existing = uow.connection.execute(
                        """
                        SELECT row_fingerprint,normalized_row_json FROM
                        infrastructure_workbook_staging_rows
                        WHERE workbook_run_id=? AND sheet_kind=? AND row_ordinal=?
                        """,
                        (run_id, sheet_kind, ordinal),
                    ).fetchone()
                    if existing is None:
                        uow.connection.execute(
                            """
                            INSERT INTO infrastructure_workbook_staging_rows(
                                staging_row_id,workbook_run_id,sheet_kind,row_ordinal,
                                row_fingerprint,normalized_row_json,validation_state,
                                warning_codes_json
                            ) VALUES (?,?,?,?,?,?,'valid','[]')
                            """,
                            (
                                new_uuid4(),
                                run_id,
                                sheet_kind,
                                ordinal,
                                fingerprint,
                                encoded,
                            ),
                        )
                    elif tuple(existing) != (fingerprint, encoded):
                        raise IntegrityFailure(
                            "Infrastructure staged row changed across durable replay"
                        )
                    cursors[sheet_kind] = ordinal
                checkpoint = {
                    **checkpoint,
                    "phase": "staging",
                    "last_committed_network_element_row": cursors["network_elements"],
                    "last_committed_ip_row": cursors["ip_addresses"],
                }
                self._jobs.checkpoint_in_uow(uow, claim, checkpoint)

        for sheet_kind, ordinal, normalized, fingerprint in self._iter_normalized(
            captured,
            same_installation=same_installation,
        ):
            cursor = cursors[sheet_kind]
            if cursor is not None and ordinal <= int(cursor):
                with ReadSnapshot(self._factory) as snapshot:
                    existing = snapshot.connection.execute(
                        """
                        SELECT row_fingerprint,normalized_row_json FROM
                        infrastructure_workbook_staging_rows
                        WHERE workbook_run_id=? AND sheet_kind=? AND row_ordinal=?
                        """,
                        (run_id, sheet_kind, ordinal),
                    ).fetchone()
                encoded = canonical_json_bytes(normalized).decode("utf-8")
                if existing is None or tuple(existing) != (fingerprint, encoded):
                    raise IntegrityFailure(
                        "Infrastructure staging cursor no longer matches persisted rows"
                    )
                continue
            batch.append((sheet_kind, ordinal, normalized, fingerprint))
            if len(batch) >= _STAGE_BATCH:
                commit(batch)
                batch = []
        commit(batch)
        return checkpoint

    def _build_proposals(
        self,
        claim: DurableJobClaim,
        checkpoint: dict,
        *,
        same_installation: bool,
    ) -> None:
        run_id = str(checkpoint["current_workbook_run_id"])
        after: tuple[str, int] | None = None
        while True:
            with ReadSnapshot(self._factory) as snapshot:
                if after is None:
                    rows = snapshot.connection.execute(
                        """
                        SELECT staging_row_id,sheet_kind,row_ordinal,row_fingerprint,
                               normalized_row_json
                        FROM infrastructure_workbook_staging_rows
                        WHERE workbook_run_id=?
                        ORDER BY sheet_kind,row_ordinal LIMIT ?
                        """,
                        (run_id, _PROPOSAL_BATCH),
                    ).fetchall()
                else:
                    rows = snapshot.connection.execute(
                        """
                        SELECT staging_row_id,sheet_kind,row_ordinal,row_fingerprint,
                               normalized_row_json
                        FROM infrastructure_workbook_staging_rows
                        WHERE workbook_run_id=?
                          AND (sheet_kind>? OR (sheet_kind=? AND row_ordinal>?))
                        ORDER BY sheet_kind,row_ordinal LIMIT ?
                        """,
                        (run_id, after[0], after[0], after[1], _PROPOSAL_BATCH),
                    ).fetchall()
                prepared = []
                for row in rows:
                    normalized = loads_canonical_json(
                        str(row[4]),
                        max_bytes=131_072,
                        max_depth=4,
                        max_collection_items=64,
                    )
                    proposal = build_workbook_proposal(
                        snapshot,
                        sheet_kind=str(row[1]),
                        normalized_row=normalized,
                        row_fingerprint=str(row[3]),
                        same_installation=same_installation,
                    )
                    prepared.append((tuple(row), proposal))
            if not rows:
                break
            with UnitOfWork(self._factory) as uow:
                self._jobs.assert_claim_current(uow, claim)
                for row, proposal in prepared:
                    staging_id, sheet_kind, ordinal, row_fingerprint, _normalized_json = row
                    existing = uow.connection.execute(
                        """
                        SELECT proposal_id,action,target_network_element_id,
                               expected_target_revision,input_fingerprint,impact_json,
                               candidate_ids_json
                        FROM infrastructure_workbook_proposals
                        WHERE staging_row_id=? ORDER BY proposal_id LIMIT 2
                        """,
                        (staging_id,),
                    ).fetchall()
                    impact_json = canonical_json_bytes(proposal["impact"]).decode("utf-8")
                    candidate_json = canonical_json_bytes(proposal["candidate_ids"]).decode("utf-8")
                    if len(existing) > 1:
                        raise IntegrityFailure(
                            "Infrastructure staging row has multiple review proposals"
                        )
                    if existing:
                        current = existing[0]
                        if tuple(current[1:]) != (
                            proposal["action"],
                            proposal["target_network_element_id"],
                            proposal["expected_target_revision"],
                            proposal["input_fingerprint"],
                            impact_json,
                            candidate_json,
                        ):
                            raise IntegrityFailure(
                                "Infrastructure workbook proposal changed across replay"
                            )
                    else:
                        uow.connection.execute(
                            """
                            INSERT INTO infrastructure_workbook_proposals(
                                proposal_id,workbook_run_id,staging_row_id,action,state,
                                target_network_element_id,expected_target_revision,
                                input_fingerprint,impact_json,created_at_utc,
                                last_command_id,revision,candidate_ids_json
                            ) VALUES (?,?,?,?,'pending',?,?,?,?,?,NULL,1,?)
                            """,
                            (
                                new_uuid4(),
                                run_id,
                                staging_id,
                                proposal["action"],
                                proposal["target_network_element_id"],
                                proposal["expected_target_revision"],
                                proposal["input_fingerprint"],
                                impact_json,
                                int(self._clock()),
                                candidate_json,
                            ),
                        )
                    warnings_json = canonical_json_bytes(
                        proposal["warning_codes"]
                    ).decode("utf-8")
                    validation_state = (
                        "invalid"
                        if proposal["action"] == "skip_invalid"
                        else "warning"
                        if proposal["warning_codes"]
                        else "valid"
                    )
                    uow.connection.execute(
                        """
                        UPDATE infrastructure_workbook_staging_rows
                        SET validation_state=?,warning_codes_json=?
                        WHERE staging_row_id=?
                        """,
                        (validation_state, warnings_json, staging_id),
                    )
            last = rows[-1]
            after = (str(last[1]), int(last[2]))

    def _publish_run(
        self,
        claim: DurableJobClaim,
        checkpoint: dict,
        *,
        directory: Path,
        candidate: WorkbookCandidate,
        file_sha256: str,
        logical_fingerprint: str,
    ) -> dict:
        run_id = str(checkpoint["current_workbook_run_id"])
        recaptured = self._capture(directory, candidate)
        try:
            summary = inspect_infrastructure_workbook(
                recaptured,
                current_data_instance_id=self._payload(claim)["data_instance_id"],
            )
            if (
                recaptured.content_sha256 != file_sha256
                or summary.logical_fingerprint != logical_fingerprint
            ):
                raise SomaError(
                    "WORKBOOK_STALE",
                    "Infrastructure workbook changed before review publication",
                )
        finally:
            recaptured.close()

        publishing = {**checkpoint, "phase": "publishing"}
        self._jobs.checkpoint(claim, publishing)
        checkpoint = publishing
        with UnitOfWork(self._factory) as uow:
            self._jobs.assert_claim_current(uow, claim)
            run = uow.connection.execute(
                """
                SELECT state,network_element_row_count,ip_row_count,revision
                FROM infrastructure_workbook_runs WHERE workbook_run_id=?
                """,
                (run_id,),
            ).fetchone()
            if run is None or str(run[0]) != "validating" or int(run[3]) != 1:
                raise IntegrityFailure("Infrastructure run is not publishable")
            staged_counts = dict(
                uow.connection.execute(
                    """
                    SELECT sheet_kind,count(*) FROM infrastructure_workbook_staging_rows
                    WHERE workbook_run_id=? GROUP BY sheet_kind
                    """,
                    (run_id,),
                ).fetchall()
            )
            if (
                int(staged_counts.get("network_elements", 0)) != int(run[1])
                or int(staged_counts.get("ip_addresses", 0)) != int(run[2])
            ):
                raise IntegrityFailure("Infrastructure staging counts are incomplete")
            proposal_count = int(
                uow.connection.execute(
                    "SELECT count(*) FROM infrastructure_workbook_proposals "
                    "WHERE workbook_run_id=?",
                    (run_id,),
                ).fetchone()[0]
            )
            if proposal_count != int(run[1]) + int(run[2]):
                raise IntegrityFailure("Infrastructure proposal set is incomplete")
            warning_count = int(
                uow.connection.execute(
                    "SELECT coalesce(sum(json_array_length(warning_codes_json)),0) "
                    "FROM infrastructure_workbook_staging_rows WHERE workbook_run_id=?",
                    (run_id,),
                ).fetchone()[0]
            )
            uow.connection.execute(
                """
                UPDATE infrastructure_workbook_runs
                SET state='staged',published_at_utc=?,warning_count=?,revision=2
                WHERE workbook_run_id=? AND state='validating' AND revision=1
                """,
                (int(self._clock()), warning_count, run_id),
            )
            published_ids = list(checkpoint["published_run_ids"])
            if run_id not in published_ids:
                if len(published_ids) >= _MAX_CANDIDATES:
                    raise IntegrityFailure("Infrastructure published-run recovery cache overflow")
                published_ids.append(run_id)
            next_checkpoint = _empty_checkpoint(
                manifest_sha256=str(checkpoint["candidate_manifest_sha256"]),
                candidate_count=int(checkpoint["candidate_count"]),
                candidate_index=int(checkpoint["candidate_index"]) + 1,
                published_run_ids=published_ids,
                phase="waiting_review",
            )
            self._jobs.checkpoint_in_uow(uow, claim, next_checkpoint)
        return next_checkpoint

    def run(self, claim: DurableJobClaim) -> InfrastructureWorkbookStageResult:
        payload = self._payload(claim)
        checkpoint = self._checkpoint(claim)
        directory, candidates, manifest = self._discover(payload)

        if checkpoint is None:
            checkpoint = _empty_checkpoint(
                manifest_sha256=manifest,
                candidate_count=len(candidates),
            )
            self._jobs.checkpoint(claim, checkpoint)
        elif (
            checkpoint["candidate_manifest_sha256"] != manifest
            or checkpoint["candidate_count"] != len(candidates)
        ):
            checkpoint = _empty_checkpoint(
                manifest_sha256=manifest,
                candidate_count=len(candidates),
                candidate_index=0,
                published_run_ids=checkpoint["published_run_ids"],
            )
            self._jobs.checkpoint(claim, checkpoint)

        while int(checkpoint["candidate_index"]) < len(candidates):
            index = int(checkpoint["candidate_index"])
            candidate = candidates[index]
            captured = self._capture(directory, candidate)
            try:
                summary = inspect_infrastructure_workbook(
                    captured,
                    current_data_instance_id=payload["data_instance_id"],
                )
                published = self._existing_review_run(
                    candidate,
                    file_sha256=captured.content_sha256,
                    logical_fingerprint=summary.logical_fingerprint,
                )
                if published is not None:
                    published_ids = list(checkpoint["published_run_ids"])
                    if published not in published_ids:
                        published_ids.append(published)
                    checkpoint = _empty_checkpoint(
                        manifest_sha256=manifest,
                        candidate_count=len(candidates),
                        candidate_index=index + 1,
                        published_run_ids=published_ids,
                    )
                    self._jobs.checkpoint(claim, checkpoint)
                    continue

                if checkpoint["current_workbook_run_id"] is None:
                    _run_id, checkpoint = self._start_run(
                        claim,
                        checkpoint,
                        candidate=candidate,
                        file_sha256=captured.content_sha256,
                        summary=summary,
                    )
                else:
                    self._require_run_identity(
                        checkpoint,
                        candidate,
                        file_sha256=captured.content_sha256,
                        summary=summary,
                    )
                checkpoint = self._stage_rows(
                    claim,
                    checkpoint,
                    captured,
                    same_installation=summary.same_installation,
                )
                self._build_proposals(
                    claim,
                    checkpoint,
                    same_installation=summary.same_installation,
                )
            finally:
                captured.close()

            checkpoint = self._publish_run(
                claim,
                checkpoint,
                directory=directory,
                candidate=candidate,
                file_sha256=str(checkpoint["current_file_sha256"]),
                logical_fingerprint=summary.logical_fingerprint,
            )
            if int(checkpoint["candidate_index"]) < len(candidates):
                checkpoint = {
                    **checkpoint,
                    "phase": "discovering",
                }
                self._jobs.checkpoint(claim, checkpoint)

        completed = {**checkpoint, "phase": "completed"}
        with UnitOfWork(self._factory) as uow:
            self._jobs.assert_claim_current(uow, claim)
            self._jobs.checkpoint_in_uow(uow, claim, completed)
            self._jobs.complete_in_uow(uow, claim)
        return InfrastructureWorkbookStageResult(
            job_id=claim.job_id,
            published_run_ids=tuple(completed["published_run_ids"]),
        )


__all__ = [
    "InfrastructureWorkbookStageResult",
    "InfrastructureWorkbookStageWorker",
    "WorkbookCandidate",
    "candidate_manifest_sha256",
    "discover_workbook_candidates",
]
