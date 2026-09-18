from __future__ import annotations

import hashlib
import json
import os
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Protocol

from openpyxl import Workbook, load_workbook

from soma.foundation.errors import IntegrityFailure, SomaError, ValidationError
from soma.foundation.identifiers import require_uuid4, utc_epoch_seconds
from soma.foundation.persistence.connections import ConnectionFactory
from soma.foundation.persistence.uow import ReadSnapshot
from soma.foundation.strict_json import canonical_json_bytes, sha256_canonical_json

from ..repositories.reports import ReportAttemptRecord, ReportRepository

_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_ALLOWED_ARTIFACT_STATES = {"ready_to_generate", "generating", "verifying", "completed"}
_SHEET_NAMES = ("_SOMA_Metadata", "Members", "Tiers", "Cohorts", "Sections")
_MEMBER_COLUMNS = (
    "member_ordinal",
    "service_request_id",
    "customer_org_id",
    "contract_id",
    "contract_product_line_id",
    "policy_revision_id",
    "classification_event_id",
    "severity",
    "report_date_utc",
    "status_class",
    "endpoint_utc",
    "suspension_num",
    "suspension_den",
    "elapsed_num",
    "elapsed_den",
    "calculation_state",
    "sla_input_token",
    "source_report_date_evidence_id",
    "source_status_evidence_id",
    "source_suspension_evidence_id",
    "display_values_json",
)
_TIER_COLUMNS = (
    "service_request_id",
    "policy_tier_id",
    "individual_state",
    "inclusive_boundary_met",
)
_COHORT_COLUMNS = (
    "cohort_ordinal",
    "calendar_month",
    "customer_org_id",
    "contract_id",
    "contract_product_line_id",
    "policy_revision_id",
    "policy_tier_id",
    "severity",
    "denominator",
    "terminal_met_count",
    "terminal_exceeded_count",
    "active_within_count",
    "active_exceeded_count",
    "state",
    "is_final",
    "input_fingerprint",
)
_SECTION_COLUMNS = (
    "section_kind",
    "schema_name",
    "schema_version",
    "section_ordinal",
    "row_ordinal",
    "canonical_row_key",
    "payload_json",
    "payload_sha256",
)


class ReportDestinationResolver(Protocol):
    def resolve_directory(self, destination_request_token: str) -> Path: ...


@dataclass(frozen=True, slots=True)
class BoundReportDestination:
    destination_request_token: str
    directory: Path

    def resolve_directory(self, destination_request_token: str) -> Path:
        if destination_request_token != self.destination_request_token:
            raise SomaError(
                "SLA_REPORT_SCOPE_INVALID",
                "report destination token does not match the validated destination context",
            )
        try:
            resolved = self.directory.resolve(strict=True)
        except OSError as exc:
            raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "report destination is unavailable") from exc
        if not resolved.is_absolute() or not resolved.is_dir():
            raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "report destination is not a local directory")
        return resolved


@dataclass(frozen=True, slots=True)
class VerifiedReportArtifact:
    report_attempt_id: str
    snapshot_hash: str
    candidate_filename: str
    artifact_sha256: str
    artifact_size_bytes: int
    verified_at_utc: int


def _filename(value: str, *, extension: str = ".xlsx") -> str:
    if (
        not isinstance(value, str)
        or not value
        or len(value.encode("utf-8")) > 4096
        or any(ch in value for ch in ("/", "\\", "\x00"))
        or value in {".", ".."}
        or not value.lower().endswith(extension)
    ):
        raise ValidationError("report artifact filename is invalid")
    return value


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while True:
            chunk = stream.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def _canonical_json_text(value: object) -> str:
    return canonical_json_bytes(value).decode("utf-8")


def _read_json_cell(value: object, *, label: str) -> object:
    if value is None:
        return None
    if not isinstance(value, str):
        raise IntegrityFailure(f"{label} must be canonical JSON text or blank")
    try:
        return json.loads(value)
    except (TypeError, ValueError) as exc:
        raise IntegrityFailure(f"{label} is invalid JSON") from exc


def _canonicalize_xlsx_archive(path: Path) -> None:
    source = path.with_name(path.name + ".zip-source")
    os.replace(path, source)
    try:
        with zipfile.ZipFile(source, "r") as reader, zipfile.ZipFile(
            path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as writer:
            for name in sorted(reader.namelist()):
                data = reader.read(name)
                info = zipfile.ZipInfo(filename=name, date_time=_FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0
                info.flag_bits = 0
                writer.writestr(info, data)
    finally:
        try:
            source.unlink()
        except FileNotFoundError:
            pass


def _require_sealed_snapshot(
    factory: ConnectionFactory,
    *,
    report_attempt_id: str,
    snapshot_hash: str,
) -> tuple[ReportAttemptRecord, dict[str, object]]:
    report_id = require_uuid4(report_attempt_id)
    if not isinstance(snapshot_hash, str) or len(snapshot_hash) != 64:
        raise ValidationError("snapshot_hash must be lowercase SHA-256")
    with ReadSnapshot(factory) as snapshot:
        attempt = ReportRepository.get_attempt(snapshot.connection, report_id)
        if attempt is None:
            raise SomaError("SLA_REPORT_ATTEMPT_STATE", "report attempt does not exist")
        if (
            attempt.state not in _ALLOWED_ARTIFACT_STATES
            or attempt.snapshot_hash != snapshot_hash
        ):
            raise SomaError("SLA_REPORT_ATTEMPT_STATE", "report attempt is not sealed for artifact use")
        semantic = ReportRepository.snapshot_semantic_value(snapshot.connection, report_id)
        actual_hash = sha256_canonical_json(semantic)
        if actual_hash != snapshot_hash:
            raise SomaError("SLA_REPORT_SNAPSHOT_MISMATCH", "sealed report snapshot hash changed")
        members, tiers, cohorts, sections = ReportRepository.counts(snapshot.connection, report_id)
        if (
            members != attempt.snapshot_member_count
            or cohorts != attempt.snapshot_cohort_count
            or sections != attempt.snapshot_section_row_count
            or tiers != len(semantic["tiers"])  # type: ignore[arg-type]
        ):
            raise SomaError("SLA_REPORT_SNAPSHOT_MISMATCH", "sealed report counts changed")
        return attempt, semantic


def default_final_filename(attempt: ReportAttemptRecord) -> str:
    scope = "all" if attempt.customer_org_id is None else "customer"
    short_id = attempt.report_attempt_id.split("-", 1)[0]
    return _filename(f"SOMA-SLA-{attempt.period_type}-{scope}-{short_id}.xlsx")


def default_candidate_filename(attempt: ReportAttemptRecord, snapshot_hash: str) -> str:
    final = default_final_filename(attempt)
    return _filename(f".{final[:-5]}.candidate-{snapshot_hash[:12]}.xlsx")


class SlaReportXlsxArtifact:
    def __init__(
        self,
        connection_factory: ConnectionFactory,
        destination_resolver: ReportDestinationResolver,
    ) -> None:
        self._factory = connection_factory
        self._destinations = destination_resolver

    @staticmethod
    def _append_table(sheet, columns: tuple[str, ...], rows: list[dict[str, object]]) -> None:
        sheet.append(columns)
        for row in rows:
            values: list[object] = []
            for column in columns:
                value = row.get(column)
                if column in {"display_values_json", "payload_json"}:
                    if value is not None and not isinstance(value, str):
                        value = _canonical_json_text(value)
                values.append(value)
            sheet.append(tuple(values))

    def write_candidate(
        self,
        *,
        report_attempt_id: str,
        snapshot_hash: str,
        destination_request_token: str,
        candidate_filename: str | None = None,
    ) -> str:
        attempt, semantic = _require_sealed_snapshot(
            self._factory,
            report_attempt_id=report_attempt_id,
            snapshot_hash=snapshot_hash,
        )
        candidate = _filename(
            candidate_filename or default_candidate_filename(attempt, snapshot_hash)
        )
        directory = self._destinations.resolve_directory(destination_request_token)
        target = directory / candidate
        if target.parent != directory:
            raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "candidate escaped validated destination")
        if target.exists():
            raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "candidate filename already exists")

        fd, temporary_name = tempfile.mkstemp(
            prefix=f".{candidate}.",
            suffix=".tmp",
            dir=directory,
        )
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            workbook = Workbook(write_only=True)
            workbook.properties.creator = "SOMA"
            workbook.properties.title = "SOMA SLA Report"
            workbook.properties.created = datetime(1980, 1, 1)
            workbook.properties.modified = datetime(1980, 1, 1)

            metadata = workbook.create_sheet("_SOMA_Metadata")
            request = semantic["report_request"]
            if not isinstance(request, dict):
                raise IntegrityFailure("sealed report request metadata is invalid")
            metadata_rows = (
                ("contract_id", "SLA_REPORT_XLSX_V1"),
                ("contract_version", 1),
                ("report_attempt_id", attempt.report_attempt_id),
                ("snapshot_hash", snapshot_hash),
                ("represented_content_sha256", sha256_canonical_json(semantic)),
                ("as_of_utc", attempt.as_of_utc),
                ("period_timezone", attempt.period_timezone),
                ("scope_fingerprint", ReportRepository.scope_fingerprint(attempt)),
                ("member_count", len(semantic["members"])),  # type: ignore[arg-type]
                ("tier_count", len(semantic["tiers"])),  # type: ignore[arg-type]
                ("cohort_count", len(semantic["cohorts"])),  # type: ignore[arg-type]
                ("section_count", len(semantic["sections"])),  # type: ignore[arg-type]
                ("report_request_json", _canonical_json_text(request)),
            )
            metadata.append(("key", "value"))
            for item in metadata_rows:
                metadata.append(item)

            members = workbook.create_sheet("Members")
            member_rows = []
            for raw in semantic["members"]:  # type: ignore[union-attr]
                row = dict(raw)
                row["display_values_json"] = (
                    None
                    if row.pop("display_values") is None
                    else _canonical_json_text(raw["display_values"])  # type: ignore[index]
                )
                member_rows.append(row)
            self._append_table(members, _MEMBER_COLUMNS, member_rows)

            tiers = workbook.create_sheet("Tiers")
            self._append_table(tiers, _TIER_COLUMNS, [dict(row) for row in semantic["tiers"]])  # type: ignore[arg-type]

            cohorts = workbook.create_sheet("Cohorts")
            self._append_table(cohorts, _COHORT_COLUMNS, [dict(row) for row in semantic["cohorts"]])  # type: ignore[arg-type]

            sections = workbook.create_sheet("Sections")
            section_rows = []
            for raw in semantic["sections"]:  # type: ignore[union-attr]
                row = dict(raw)
                row["payload_json"] = _canonical_json_text(row.pop("payload"))
                section_rows.append(row)
            self._append_table(sections, _SECTION_COLUMNS, section_rows)

            workbook.save(temporary)
            _canonicalize_xlsx_archive(temporary)
            if temporary.stat().st_size <= 0:
                raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "candidate workbook is empty")
            os.replace(temporary, target)
        except SomaError:
            try:
                temporary.unlink()
            except FileNotFoundError:
                pass
            raise
        except BaseException as exc:
            try:
                temporary.unlink()
            except FileNotFoundError:
                pass
            raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "could not write report workbook") from exc
        return candidate

    @staticmethod
    def _metadata(workbook) -> dict[str, object]:
        sheet = workbook["_SOMA_Metadata"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header != ("key", "value"):
            raise IntegrityFailure("report metadata header is invalid")
        result: dict[str, object] = {}
        for row in rows:
            if len(row) < 2 or not isinstance(row[0], str) or row[0] in result:
                raise IntegrityFailure("report metadata row is invalid")
            result[row[0]] = row[1]
        return result

    @staticmethod
    def _read_table(workbook, name: str, columns: tuple[str, ...]) -> list[dict[str, object]]:
        sheet = workbook[name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header != columns:
            raise IntegrityFailure(f"{name} header is invalid")
        output: list[dict[str, object]] = []
        for raw in rows:
            if len(raw) != len(columns):
                raise IntegrityFailure(f"{name} row width is invalid")
            output.append(dict(zip(columns, raw, strict=True)))
        return output

    def verify_candidate(
        self,
        *,
        report_attempt_id: str,
        snapshot_hash: str,
        destination_request_token: str,
        candidate_filename: str,
    ) -> VerifiedReportArtifact:
        attempt, _semantic = _require_sealed_snapshot(
            self._factory,
            report_attempt_id=report_attempt_id,
            snapshot_hash=snapshot_hash,
        )
        candidate = _filename(candidate_filename)
        directory = self._destinations.resolve_directory(destination_request_token)
        path = directory / candidate
        try:
            if not path.is_file() or path.stat().st_size <= 0:
                raise IntegrityFailure("report candidate is missing or empty")
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                if tuple(workbook.sheetnames) != _SHEET_NAMES:
                    raise IntegrityFailure("report workbook sheet contract is invalid")
                metadata = self._metadata(workbook)
                expected_metadata = {
                    "contract_id": "SLA_REPORT_XLSX_V1",
                    "contract_version": 1,
                    "report_attempt_id": attempt.report_attempt_id,
                    "snapshot_hash": snapshot_hash,
                    "represented_content_sha256": snapshot_hash,
                    "as_of_utc": attempt.as_of_utc,
                    "period_timezone": attempt.period_timezone,
                    "scope_fingerprint": ReportRepository.scope_fingerprint(attempt),
                }
                for key, expected in expected_metadata.items():
                    if metadata.get(key) != expected:
                        raise IntegrityFailure(f"report workbook metadata mismatch: {key}")

                members = self._read_table(workbook, "Members", _MEMBER_COLUMNS)
                tiers = self._read_table(workbook, "Tiers", _TIER_COLUMNS)
                cohorts = self._read_table(workbook, "Cohorts", _COHORT_COLUMNS)
                sections = self._read_table(workbook, "Sections", _SECTION_COLUMNS)
                if (
                    metadata.get("member_count") != len(members)
                    or metadata.get("tier_count") != len(tiers)
                    or metadata.get("cohort_count") != len(cohorts)
                    or metadata.get("section_count") != len(sections)
                    or len(members) != attempt.snapshot_member_count
                    or len(cohorts) != attempt.snapshot_cohort_count
                    or len(sections) != attempt.snapshot_section_row_count
                ):
                    raise IntegrityFailure("report workbook represented counts mismatch")

                report_request = _read_json_cell(
                    metadata.get("report_request_json"),
                    label="report_request_json",
                )
                if not isinstance(report_request, dict):
                    raise IntegrityFailure("report request metadata is invalid")

                reconstructed_members: list[dict[str, object]] = []
                for row in members:
                    display = _read_json_cell(row.pop("display_values_json"), label="display_values_json")
                    row["display_values"] = display
                    reconstructed_members.append(row)
                reconstructed_sections: list[dict[str, object]] = []
                for row in sections:
                    payload = _read_json_cell(row.pop("payload_json"), label="payload_json")
                    if not isinstance(payload, dict):
                        raise IntegrityFailure("report section payload is invalid")
                    row["payload"] = payload
                    reconstructed_sections.append(row)

                semantic = {
                    "schema": "SOMA_REPORT_SNAPSHOT_V1",
                    "report_request": report_request,
                    "members": reconstructed_members,
                    "tiers": tiers,
                    "cohorts": cohorts,
                    "sections": reconstructed_sections,
                }
                represented_hash = sha256_canonical_json(semantic)
                if represented_hash != snapshot_hash:
                    raise IntegrityFailure("report workbook represented content does not match sealed snapshot")
            finally:
                workbook.close()
        except (IntegrityFailure, ValidationError):
            raise
        except BaseException as exc:
            raise SomaError("SLA_REPORT_ARTIFACT_VERIFY_FAILED", "candidate workbook could not be verified") from exc

        return VerifiedReportArtifact(
            report_attempt_id=attempt.report_attempt_id,
            snapshot_hash=snapshot_hash,
            candidate_filename=candidate,
            artifact_sha256=_sha256_file(path),
            artifact_size_bytes=path.stat().st_size,
            verified_at_utc=utc_epoch_seconds(),
        )

    def publish_verified(
        self,
        *,
        verification: VerifiedReportArtifact,
        destination_request_token: str,
        final_filename: str | None = None,
    ) -> dict[str, object]:
        directory = self._destinations.resolve_directory(destination_request_token)
        candidate = directory / _filename(verification.candidate_filename)
        with ReadSnapshot(self._factory) as snapshot:
            attempt = ReportRepository.get_attempt(snapshot.connection, verification.report_attempt_id)
            if attempt is None or attempt.snapshot_hash != verification.snapshot_hash:
                raise SomaError("SLA_REPORT_ATTEMPT_STATE", "report attempt changed before publication")
            final = _filename(final_filename or default_final_filename(attempt))
        destination = directory / final
        if destination.exists():
            if (
                destination.is_file()
                and _sha256_file(destination) == verification.artifact_sha256
                and destination.stat().st_size == verification.artifact_size_bytes
            ):
                # Recovery after atomic publication but before database completion.
                pass
            else:
                raise SomaError(
                    "SLA_REPORT_ARTIFACT_WRITE_FAILED",
                    "final report filename collides with unrelated existing content",
                )
        else:
            if not candidate.is_file():
                raise SomaError("SLA_REPORT_ARTIFACT_WRITE_FAILED", "verified candidate disappeared")
            if (
                _sha256_file(candidate) != verification.artifact_sha256
                or candidate.stat().st_size != verification.artifact_size_bytes
            ):
                raise SomaError("SLA_REPORT_ARTIFACT_VERIFY_FAILED", "verified candidate bytes changed")
            os.replace(candidate, destination)

        if (
            not destination.is_file()
            or _sha256_file(destination) != verification.artifact_sha256
            or destination.stat().st_size != verification.artifact_size_bytes
        ):
            raise SomaError("SLA_REPORT_ARTIFACT_VERIFY_FAILED", "published report bytes differ from verification")
        return {
            "report_attempt_id": verification.report_attempt_id,
            "snapshot_hash": verification.snapshot_hash,
            "candidate_filename": verification.candidate_filename,
            "artifact_filename": final,
            "artifact_sha256": verification.artifact_sha256,
            "artifact_size_bytes": verification.artifact_size_bytes,
            "verified_at_utc": verification.verified_at_utc,
        }


__all__ = [
    "BoundReportDestination",
    "ReportDestinationResolver",
    "SlaReportXlsxArtifact",
    "VerifiedReportArtifact",
    "default_candidate_filename",
    "default_final_filename",
]
